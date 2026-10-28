"""The analysis PDF and workbook mark only the notes that question a figure (D-061; CLAUDE.md rule 12; MPA WP3-B).

frontend/tests/unit/dataNoteMarking.spec.tsx holds the page and its chart PNG to the rule; this is the exports' half,
over the same shared notes fixture (`tests/fixtures/analysis_data_notes.json`) and through the real builders.

1. Marks. Every case renders as an analysis PDF (its citations cited in the narrative) and a workbook. Every letter the
   PDF draws (a raised, linked letter in the grid, the methods, the citations, the narrative or the summary line), links
   to or sets in the detail's letter column, and every letter the workbook sets in a number format, names in a comment
   or the Sources Note column, or prints in the Overview's fine print or the Metrics key, is one of the case's
   `marked_letters`; and each of them shows: the PDF draws the citation supplements' too, the workbook, which carries no
   narrative, the report's. A site that bypasses `marks_figure` draws record_only_and_mixed_marking's record-only c, its
   unknown d or its gap's e. A format the reader does not know fails, so a new way of setting a letter cannot hide.
2. Every note stays: each keeps its anchored entry in the PDF (the report's under the grid, a supplement's under
   Narrative sources) and its words in the workbook's Overview, and an annual PDF prints no †.
3. The fine print is the page's (DataNotes.tsx): its places, its counts, and "No figures need review" held back beside a
   record-only note on a growth figure alone, on dataNoteMarking.spec.tsx's own report.
4. The calculated sentence names the metrics metrics-table.spec.tsx's sentences name, read from that spec, for the same
   runs.
5. The workbook's letter: a marked figure's cell stays a number whose format prints the letter, in every unit, a
   window's change and the top line's growth among them, and "n/m" takes it from the text section; a chart's value axis
   keeps the figures' own format, unlinked, so no tick prints a letter. A quarterly PDF's Q4 header takes the page's †
   where a fourth quarter's flow figure is calculated, with its legend, and no figure takes one.
6. The Metrics header row ends in the key to the letters the grid prints, exactly when it prints one.
7. The exports' own sentences hold the reason words' surface rules (test_analysis_data_notes.py): no page action, no
   drawing, no single figure, no semicolon. The gap words alone say how each file draws a gap, which is their job.
8. The PDF's summary line stays on the grid's last page (WeasyPrint): at the row count where, without the rule, it
   would start the next page, the last row moves with it; the sentence above the grid shares the grid's first page.
"""
from __future__ import annotations

import io
import json
import re
import zipfile
from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.services.analysis_export_metadata import (
    DERIVED_Q4_LEGEND, PDF_CLOSING, QUARTER_LEGEND, WORKBOOK_CLOSING, calculated_sentence, detail_entries,
    fine_print_summary, summary_places,
)
from app.services.excel_export_service import build_analysis_workbook
from app.services.export_service import ANALYSIS_GRID_PRECISION, ANALYSIS_GRID_ROUNDING, ExportService
from app.services.trend_analysis.data_notes import data_notes, marks_figure
from tests.unit.test_analysis_data_notes import DRAWING, _surface_problems

FIXTURE = json.loads((Path(__file__).parents[1] / "fixtures" / "analysis_data_notes.json").read_text())
CASES = FIXTURE["cases"]
COMPANY = SimpleNamespace(name="Test Co", ticker="TST")
FRONTEND = Path(__file__).resolve().parents[3] / "frontend"

PDF_MARK = re.compile(r'<sup class="data-note"><a href="#analysis-data-note-([a-z]+)">([a-z]+)</a></sup>')
PDF_LIST_LETTER = re.compile(r'<span class="note-letter">([a-z]+)\.</span>')
PDF_LINK = re.compile(r'href="#analysis-data-note-([a-z]+)"')
# Every format the workbook sets, with the letter a marked note adds: '#,##0" a"', '0.00"x a"', '0.0" pp a"', '@" a"'.
FORMAT = re.compile(r'(?:#,##0(?:\.00)?|0\.0%|0\.00"x|0\.0" pp|@|General)(?:"? (?P<letter>[a-z]+))?"?')


def _analysis(case: dict) -> SimpleNamespace:
    dataset = deepcopy(case["dataset"])
    dataset["data_notes"] = deepcopy(case["expected"])
    citations = [{**citation, "n": index, "excerpt": f"Source {index}"}
                 for index, citation in enumerate(deepcopy(case.get("citations", [])), start=1)]
    if "expected_citation_notes" in case:
        dataset["narrative_presentation"] = {"notes": deepcopy(case["expected_citation_notes"])}
    narrative = " ".join(f"Figure [{citation['n']}]." for citation in citations) or "No figures cited."
    return SimpleNamespace(mode=dataset["mode"], period_key=dataset["period_key"], dataset_json=dataset,
                           narrative_md=narrative, citations_json=citations)


def pdf_letters(html: str) -> set[str]:
    marks = PDF_MARK.findall(html)
    assert all(target == shown for target, shown in marks), "a letter links to another note's entry"
    return {shown for _, shown in marks} | set(PDF_LIST_LETTER.findall(html)) | set(PDF_LINK.findall(html))


def workbook_letters(data: bytes) -> set[str]:
    book = load_workbook(io.BytesIO(data))
    letters: set[str] = set()
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                found = FORMAT.fullmatch(cell.number_format)
                assert found, f"{sheet.title}!{cell.coordinate}: a format the reader does not know, {cell.number_format!r}"
                letters |= {found["letter"]} - {None}
                if cell.comment is not None:
                    letters |= set(re.findall(r"^Data note ([a-z]+):", cell.comment.text))
                if sheet.title == "Sources & methods" and cell.column == 17 and isinstance(cell.value, str):
                    letters |= set(re.findall(r"^Data note ([a-z]+)$", cell.value))
    overview = [str(row[0]) for row in book["Overview"].iter_rows(values_only=True) if row[0]]
    for text in overview[overview.index("Data notes"):]:
        # The summary's "a Revenue FY2024: …" clauses and the detail's "a. …" entries.
        letters |= set(re.findall(r"(?:^| · )([a-z]+) ", text)) | set(re.findall(r"^([a-z]+)\. ", text))
    key = [cell.value for cell in book["Metrics"][1] if cell.value and cell.fill.fill_type is None]
    letters |= {letter for text in key for letter in re.findall(r" · ([a-z]+) ", text)}
    return letters


# A case whose reason codes hold a lone surrogate, which no XML document can carry: the workbook's Sources sheet prints a
# value's reason codes (as before WP3-B), so only its PDF is read.
NO_WORKBOOK = {"portable_reason_identity_utf16_ascii": "its reason codes hold a lone surrogate (U+D800)"}


@pytest.mark.parametrize("case", CASES, ids=[case["name"] for case in CASES])
def test_the_exports_draw_exactly_the_letters_of_notes_that_question_a_figure(case):
    analysis = _analysis(case)
    html = ExportService().generate_analysis_pdf_html(analysis, COMPANY)
    assert pdf_letters(html) == set(case["marked_letters"])
    # Every note stays: its anchored entry in the PDF.
    notes = (case.get("expected_citation_notes") or case["expected"])["notes"]
    assert [note["id"] for note in notes if f'<li id="{note["id"]}">' not in html] == []
    if case["dataset"]["mode"] == "annual":
        assert "†" not in html
    if case["name"] in NO_WORKBOOK:
        return
    workbook = build_analysis_workbook(deepcopy(analysis.dataset_json))
    assert workbook_letters(workbook) == {note["letter"] for note in case["expected"]["notes"] if marks_figure(note)}
    # And its words in the workbook's fine print.
    overview = "\n".join(str(row[0]) for row in load_workbook(io.BytesIO(workbook))["Overview"].iter_rows(values_only=True)
                         if row[0])
    assert [note["letter"] for note in case["expected"]["notes"] if note["message"] not in overview] == []


def test_only_the_named_case_skips_the_workbook():
    case = next(case for case in CASES if case["name"] in NO_WORKBOOK)
    with pytest.raises(ValueError):  # lxml refuses the surrogate
        build_analysis_workbook(deepcopy({**case["dataset"], "data_notes": case["expected"]}))


def test_the_gate_reads_a_letter_wherever_an_export_can_set_one():
    """Planted: the record-only case's c, unknown d and gap e set where a site that bypasses the rule would set them."""
    case = next(case for case in CASES if case["name"] == "record_only_and_mixed_marking")
    html = ExportService().generate_analysis_pdf_html(_analysis(case), COMPANY)
    for planted, letter in (('$1<sup class="data-note"><a href="#analysis-data-note-c">c</a></sup>', "c"),
                            ('<span class="note-letter">d.</span>', "d"), ('<a href="#analysis-data-note-e">', "e")):
        assert pdf_letters(html + planted) - pdf_letters(html) == {letter}, planted

    def letters_after(plant) -> set[str]:
        book = load_workbook(io.BytesIO(build_analysis_workbook(deepcopy(case["dataset"]))))
        plant(book)
        buffer = io.BytesIO()
        book.save(buffer)
        return workbook_letters(buffer.getvalue()) - {"a", "b"}

    def formats(book):
        for cell, number_format in (("D2", '#,##0" c"'), ("D3", '0.00"x d"'), ("E2", '0.0" pp e"'), ("E3", '@" e"')):
            book["Metrics"][cell].number_format = number_format

    assert letters_after(formats) == {"c", "d", "e"}
    assert letters_after(lambda book: setattr(book["Metrics"]["D5"].comment, "text", "Data note e: Not in our data.")) == {"e"}
    assert letters_after(lambda book: book["Overview"].append(["c. Shown as stored."])) == {"c"}
    with pytest.raises(AssertionError, match="a format the reader does not know"):
        letters_after(lambda book: setattr(book["Metrics"]["D2"], "number_format", '#,##0"(c)"'))


# ---- The fine print, as the page's DataNotes.tsx ---------------------------------------------------------------
def _growth_only(yoy: dict) -> dict:
    """dataNoteMarking.spec.tsx's report: every value passed and none lacks its record; the FY2024 YoY carries `yoy`."""
    passed = {"version": 1, "method": "reported", "validation": "passed", "reasons": [], "formula": None, "inputs": []}
    dataset = {"mode": "annual", "period_key": "FY2023..FY2024", "periods": [{"key": "FY2023"}, {"key": "FY2024"}],
               "series": [{"concept": "revenue", "label": "Revenue", "unit": "USD", "points": [
                   {"period": "FY2023", "value": 100, "provenance": passed},
                   {"period": "FY2024", "value": 110, "provenance": passed, "yoy": 0.1, "yoy_provenance": yoy}]}]}
    return data_notes(dataset)


def test_the_summary_never_clears_a_report_beside_a_record_only_note():
    legacy = _growth_only({"method": "calculated", "validation": "needs_review", "reasons": ["legacy_calculation"]})
    # Every value passed and none is unknown: the counts alone would clear the report.
    assert {key: legacy["totals"][key] for key in ("reviewValues", "reviewComparisons", "unknown", "passed")} == {
        "reviewValues": 0, "reviewComparisons": 1, "unknown": 0, "passed": 2}
    assert [clause.words for clause in fine_print_summary(legacy)] == []
    cleared = _growth_only({"method": "calculated", "validation": "passed", "reasons": []})
    assert [clause.words for clause in fine_print_summary(cleared)] == ["No figures need review · 2 checked"]


def test_the_summary_and_detail_read_as_the_page_reads_them():
    case = next(case for case in CASES if case["name"] == "record_only_and_mixed_marking")
    clauses = fine_print_summary(case["expected"])
    # dataNoteMarking.spec.tsx: "a Free cash flow FY2024: not yet matched to the filing", never the record-only
    # sentence the mixed note's message opens with; then the counts.
    assert [(clause.note and clause.note["letter"], clause.words, clause.gap) for clause in clauses] == [
        ("a", "Free cash flow FY2024: not yet matched to the filing", False),
        ("b", "Long-term debt FY2024: the filings report different values (the most recent shown)", False),
        (None, "1 not in our data", True),
        (None, "1 figure with no check recorded", False),
        (None, "4 calculated from filed figures", False),
    ]
    entries = {entry["id"]: entry for entry in detail_entries(case["expected"]["notes"])}
    assert [entry["letter"] for entry in entries.values()] == ["a", "b", None, None, None]
    legacy = entries["analysis-data-note-c"]
    assert (legacy["places"], legacy["counts"]) == ("Operating margin (FY2023, FY2024).", "2 figures; 1 growth figure built on them.")
    # A note on growth figures alone names them, and the summary names at most three metrics, then counts the rest.
    note = {"values": [], "comparisons": [
        {"label": label, "period": period, "basis": basis} for label, period, basis in (
            ("Revenue", "FY2024", "yoy"), ("Revenue", "FY2022..FY2024", "cagr"), ("Net margin", "FY2022..FY2024", "window_pp"),
            ("Free cash flow", "2026Q2", "qoq"), ("Free cash flow", "2026Q3", "qoq"))]}
    assert summary_places(note) == ("Revenue YoY FY2024; Revenue CAGR FY2022–FY2024; "
                                    "Net margin window percentage-point change FY2022–FY2024 and 1 more")
    note["comparisons"] = note["comparisons"][3:]
    assert summary_places(note) == "Free cash flow QoQ Q2 2026, Q3 2026"
    assert detail_entries([{**case["expected"]["notes"][2], "values": [], "comparisons": note["comparisons"]}])[0][
        "counts"] == "2 growth figures."


# ---- The calculated sentence, as reportPresentation.ts ---------------------------------------------------------
def _run(*rows: tuple[str, str, bool, bool]) -> dict:
    """metrics-table.spec.tsx's runs: each row (concept, label, calculated, percent) with a value in both years."""
    def provenance(calculated: bool) -> dict:
        return {"version": 1, "method": "calculated" if calculated else "reported", "validation": "passed",
                "reasons": [], "formula": None, "inputs": []}
    return {"mode": "annual", "period_key": "FY2023..FY2024", "periods": [{"key": "FY2023"}, {"key": "FY2024"}],
            "series": [{"concept": concept, "label": label, "percent": percent, "points": [
                {"period": period, "value": 10 + index, "provenance": provenance(calculated)}
                for index, period in enumerate(("FY2023", "FY2024"))]} for concept, label, calculated, percent in rows]}


def test_the_calculated_sentence_names_what_the_page_names():
    revenue, fcf = ("revenue", "Revenue", False, False), ("free_cash_flow", "Free cash flow", True, False)
    operating, net = ("operating_margin", "Operating margin", True, True), ("net_margin", "Net margin", True, True)
    regn = [revenue, operating, net, fcf, ("working_capital", "Working capital", True, False),
            ("current_ratio", "Current ratio", True, False)]
    partly = _run(revenue, fcf, operating)
    partly["series"][2]["points"][1]["provenance"]["method"] = "reported"
    runs = [
        _run(revenue, operating, fcf, ("eps_diluted", "EPS (diluted)", True, False)),
        partly,
        _run(*regn),
        _run(regn[0], ("gross_margin", "Gross margin", False, True), *regn[1:]),
        _run(*regn[:3]),
    ]
    spec = (FRONTEND / "tests" / "unit" / "metrics-table.spec.tsx").read_text(encoding="utf-8")
    page = re.findall(r"'Figures are rounded for display; (.+?) calculated from filed figures\. Select any value", spec)
    assert len(page) == len(runs), "metrics-table.spec.tsx's calculated sentences moved: re-read them"
    assert [calculated_sentence(run) for run in runs] == [f"{words[:1].upper()}{words[1:]} calculated from filed figures."
                                                          for words in page]
    assert calculated_sentence(_run(revenue)) == ""


# ---- The workbook's letter ---------------------------------------------------------------------------------------
def test_a_marked_figure_stays_a_number_whose_format_prints_its_letter():
    from tests.unit.test_excel_export import _dataset

    dataset = _dataset()
    conflict = {"method": "reported", "validation": "needs_review", "reasons": ["source_conflict"]}
    revenue, margin, eps, ratio = dataset["series"][:4]
    revenue["points"][0]["provenance"] = conflict
    revenue["points"][1]["yoy_provenance"] = revenue["points"][2]["yoy_provenance"] = conflict
    revenue["cagr_provenance"] = margin["points"][0]["provenance"] = margin["window_pp_provenance"] = conflict
    eps["points"][0]["provenance"] = ratio["points"][0]["provenance"] = conflict
    data = build_analysis_workbook(dataset)
    book = load_workbook(io.BytesIO(data))
    metrics, top = book["Metrics"], book["Revenue & growth"]
    expected = {
        (metrics, "D2"): '#,##0" a"', (metrics, "D3"): '0.0%" a"', (metrics, "D4"): '#,##0.00" a"',
        (metrics, "D5"): '0.00"x a"', (metrics, "G2"): '0.0%" a"', (metrics, "G3"): '0.0" pp a"',
        (top, "B2"): '#,##0" a"', (top, "C3"): '0.0%" a"', (book["Margins"], "B2"): '0.0%" a"',
    }
    for (sheet, cell), number_format in expected.items():
        assert (sheet[cell].data_type, sheet[cell].number_format) == ("n", number_format), f"{sheet.title}!{cell}"
    # "n/m" is text: its letter comes from the format's text section.
    assert (top["C4"].value, top["C4"].data_type, top["C4"].number_format) == ("n/m", "s", '@" a"')
    # An unmarked figure keeps its format; a gap marks nothing.
    assert (metrics["E2"].number_format, metrics["E3"].number_format, metrics["D6"].number_format) == ("#,##0", "0.0%", "#,##0")
    # A chart's value axis keeps the figures' own format, unlinked, so no tick prints the letter (LibreOffice set
    # "1,600 a" on every tick of a linked axis).
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        axes = [re.findall(r'<numFmt formatCode="([^"]*)" sourceLinked="0"/>', archive.read(name).decode())
                for name in sorted(archive.namelist()) if name.startswith("xl/charts/chart")]
    assert axes == [["#,##0"], ["0.0%"], ["#,##0"], ["#,##0"]]


def test_the_metrics_header_row_keys_the_letters_its_grid_prints():
    from tests.unit.test_excel_export import _dataset

    plain = load_workbook(io.BytesIO(build_analysis_workbook(_dataset())))["Metrics"]
    assert [cell.value for cell in plain[1]][-1] == "Window"  # no marked note: no key
    dataset = _dataset()
    dataset["series"][0]["points"][0]["provenance"] = {"method": "reported", "validation": "needs_review",
                                                       "reasons": ["source_check_needed"]}
    metrics = load_workbook(io.BytesIO(build_analysis_workbook(dataset)))["Metrics"]
    headers = [cell.value for cell in metrics[1]]
    key = metrics.cell(1, headers.index("Window") + 2)
    assert key.value == "Data notes · a Revenue FY2022: not yet matched to the filing"
    assert key.fill.fill_type is None and key.border.bottom.style is None and key.alignment.wrap_text
    assert str(key.font.color.rgb).endswith("636A77") and key.font.size == 9  # the footnote font
    assert metrics.freeze_panes == "B2" and metrics["A1"].value == "Metric"


# ---- The exports' own words ----------------------------------------------------------------------------------------
def test_the_exports_own_sentences_hold_the_reason_words_surface_rules():
    regn = _run(("revenue", "Revenue", False, False), ("operating_margin", "Operating margin", True, True),
                ("net_margin", "Net margin", True, True), ("current_ratio", "Current ratio", True, False))
    case = next(case for case in CASES if case["name"] == "mixed_quarterly_quality_and_gaps")
    counts = [clause for clause in fine_print_summary(case["expected"]) if not clause.note]
    counts += fine_print_summary(_growth_only({"method": "calculated", "validation": "passed", "reasons": []}))
    assert {clause.gap for clause in counts} == {True, False} and any("growth figure" in clause.words for clause in counts)
    sentences = [ANALYSIS_GRID_ROUNDING, calculated_sentence(regn), ANALYSIS_GRID_PRECISION, PDF_CLOSING,
                 WORKBOOK_CLOSING, QUARTER_LEGEND, DERIVED_Q4_LEGEND, "Data notes for narrative sources",
                 *(clause.words for clause in counts)]
    assert {sentence: problems for sentence in sentences if (problems := _surface_problems(sentence))} == {}
    # The gap words say how each file draws a gap, so they alone draw: the PDF's dash, the workbook's empty cell.
    html = ExportService().generate_analysis_pdf_html(_analysis(case), COMPANY)
    assert ', shown as <span class="data">—</span>' in html
    overview = [row[0] for row in load_workbook(io.BytesIO(build_analysis_workbook(deepcopy(case["dataset"]))))[
        "Overview"].iter_rows(values_only=True) if row[0]]
    assert any(", left blank · " in str(text) for text in overview)
    assert DRAWING.search("shown as —") and DRAWING.search("left blank")
    # Planted: the page's own closing line and sentence ending fail here.
    assert _surface_problems("Select any figure to see its source filing, method and inputs.")
    assert _surface_problems("Select any value for its source and full precision.")


def test_a_quarterly_grid_marks_its_derived_fourth_quarters_never_a_figure():
    """MetricsTable.tsx's Q4 header † and its legend (coherence §2.5: they mark a period, not a figure): a calculated
    Q4 flow figure puts the † on its column's header, which the legend explains; no figure carries one."""
    def quarterly(derived: bool) -> SimpleNamespace:
        dataset = {"mode": "quarterly", "period_key": "2024Q3..2025Q1", "periods": [
            {"key": "2024Q3", "fiscal_period": "Q3"}, {"key": "2024Q4", "fiscal_period": "Q4"},
            {"key": "2025Q1", "fiscal_period": "Q1"}], "series": [
            {"concept": "revenue", "label": "Revenue", "unit": "USD", "points": [
                {"period": "2024Q3", "value": 300.0}, {"period": "2024Q4", "value": 400.0, "derived": derived},
                {"period": "2025Q1", "value": 350.0}]},
            {"concept": "cash_and_equivalents", "label": "Cash", "unit": "USD", "points": [{"period": "2024Q4", "value": 9.0}]}]}
        dataset["data_notes"] = data_notes(dataset)
        return SimpleNamespace(mode="quarterly", period_key=dataset["period_key"], dataset_json=dataset, narrative_md="",
                               citations_json=[])

    html = ExportService().generate_analysis_pdf_html(quarterly(True), COMPANY)
    assert '<th class="num">Q4 2024 †</th>' in html and '<th class="num">Q3 2024</th>' in html
    assert f"{QUARTER_LEGEND} {DERIVED_Q4_LEGEND}" in html
    assert html.count("†") == 2, "the header and its legend, never a figure"
    plain = ExportService().generate_analysis_pdf_html(quarterly(False), COMPANY)
    assert "†" not in plain and QUARTER_LEGEND in plain


# ---- The PDF's layout ---------------------------------------------------------------------------------------------
def _rows(count: int) -> SimpleNamespace:
    conflict = {"method": "reported", "validation": "needs_review", "reasons": ["source_conflict"]}
    passed = {"method": "reported", "validation": "passed", "reasons": []}
    dataset = {"mode": "annual", "period_key": "FY2022..FY2024", "periods": [{"key": key} for key in ("FY2022", "FY2023", "FY2024")],
               "series": [{"concept": f"metric_{index}", "label": f"Metric {index}", "unit": "USD", "points": [
                   {"period": period, "value": 1000.0 + index, "provenance": conflict if (index, period) == (0, "FY2023") else passed}
                   for period in ("FY2022", "FY2023", "FY2024")]} for index in range(count)]}
    dataset["data_notes"] = data_notes(dataset)
    return SimpleNamespace(mode="annual", period_key=dataset["period_key"], dataset_json=dataset, narrative_md="",
                           citations_json=[])


def _pages(weasyprint, html: str, last_row: str) -> dict[str, int]:
    """The page each block starts on: the sentence above the grid, the grid's first and last rows and the summary."""
    found: dict[str, int] = {}
    for number, page in enumerate(weasyprint.HTML(string=html).render().pages):
        for box in page._page_box.descendants():
            element = box.element
            if element is None:
                continue
            text, classes = "".join(element.itertext()), element.get("class") or ""
            if element.tag == "tr" and text.startswith("Metric 0"):
                found.setdefault("first row", number)
            if element.tag == "tr" and text.startswith(last_row):
                found.setdefault("last row", number)
            if element.tag == "p" and classes == "fine-print" and text.startswith(ANALYSIS_GRID_ROUNDING):
                found.setdefault("sentence", number)
            if "data-notes-summary" in classes:
                found.setdefault("summary", number)
    return found


def test_the_summary_line_stays_on_the_grids_last_page():
    weasyprint = pytest.importorskip("weasyprint")
    rule = ".data-notes-summary { break-before: avoid; } "
    for count in range(8, 60):
        html = ExportService().generate_analysis_pdf_html(_rows(count), COMPANY)
        # The metrics section alone: the layout under test, without laying out the rest of the report.
        html = re.sub(r'(<section class="metrics-landscape">.*?</section>).*(</body>)', r"\1\2", html, count=1, flags=re.S)
        # Without the rule as written here, the summary never parts from the grid (the rule spelt another way keeps
        # it there, and the loop fails), or the rule is gone, and the layout below fails at the boundary.
        try:
            without = _pages(weasyprint, html.replace(rule, ""), f"Metric {count - 1}")
        except Exception as exc:  # missing system libs in odd envs: not what this test reads
            pytest.skip(f"WeasyPrint runtime unavailable: {exc}")
        if without["summary"] != without["last row"]:
            break
    else:
        pytest.fail("no row count from 8 to 59 puts the summary on the page after the grid without the rule")
    pages = _pages(weasyprint, html, f"Metric {count - 1}")
    assert pages["summary"] == pages["last row"], (count, pages)
    assert pages["sentence"] == pages["first row"], (count, pages)
