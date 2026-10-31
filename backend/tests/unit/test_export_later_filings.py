"""The analysis PDF and workbook read the later-filing evidence (MPA WP5a-X; D9, D10; CLAUDE.md rule 12).

The page reads it since WP5a-F1 (frontend/features/analysis/lib/laterFilings.ts); the producer (WP5a-B) sends it. These
run every case of the shared fixture backend/tests/fixtures/analysis_later_filings.json, which the page's
later-filings.spec.tsx reads too, through the real builders:

1. The line. Each file prints the later-filing line exactly where the case's expected places hold a figure: one entry
   after the notes, with no letter and no anchor, naming those figures and, where growth reads n/m across the change,
   those growths, in the fixture's order. Its words are read from the page in test_export_note_marks.py.
2. The evidence. Every figure and input of the window that carries its own annual report's value gives the dialog's
   sentence in both files, with any finding: in the PDF under its Sources & methods row, with that report's accession
   and link; in the workbook in the comment of each cell that shows the figure and in a "First reported" row of Sources
   & methods. Every growth that reads n/m across the change gives its cause, then its growth as first reported, none
   where the own reports give none: in both Sources & methods and in the top line's growth comment.
3. The sentences. The closing line's evidence sentence and the workbook's comment clause print exactly where the line
   does; Sources & methods' n/m sentence exactly where some growth reads n/m across the change.
4. No letter. Evidence marks nothing: each file draws exactly the case's marked letters.
5. Parts the payload lacks are left out, never invented.
6. Inert until the producer: a dataset without the fields prints none of it. (The byte comparison of both files
   against main's, for datasets without the fields, is the pull request's.)
"""
from __future__ import annotations

import html as html_lib
import io
import json
import re
from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.services.analysis_export_metadata import (
    BASIS_CHANGE_CAUSE, LATER_FILING_EVIDENCE, LATER_FILING_MESSAGE, NOT_MEANINGFUL_CAUSES, PDF_CLOSING, WORKBOOK_CLOSING,
    disagreement_sentence, evidence_lines, first_reported_sentence, growth_text, later_filing_entry, pdf_closing,
    workbook_closing,
)
from app.services.excel_export_service import _LINE_PANELS, build_analysis_workbook
from app.services.export_service import ExportService
from app.services.trend_analysis.data_notes import data_notes, marks_figure
from tests.unit.test_export_note_marks import pdf_letters, workbook_letters

FIXTURE = Path(__file__).resolve().parents[1] / "fixtures" / "analysis_later_filings.json"
ANALYSIS = Path(__file__).resolve().parents[3] / "frontend" / "features" / "analysis"
CASES = json.loads(FIXTURE.read_text())["cases"]
IDS = [case["name"] for case in CASES]
COMPANY = SimpleNamespace(name="Regeneron Pharmaceuticals", ticker="REGN")
EVIDENCE_KEYS = ("first_reported", "yoy_first_reported")


def _case(name: str) -> dict:
    return deepcopy(next(case for case in CASES if case["name"] == name))


def _point(dataset: dict, concept: str, period: str) -> dict:
    series = next(series for series in dataset["series"] if series["concept"] == concept)
    return next(point for point in series["points"] if point["period"] == period)


def _pdf(dataset: dict) -> str:
    analysis = SimpleNamespace(mode=dataset["mode"], period_key=dataset["period_key"], dataset_json=deepcopy(dataset),
                               narrative_md="", citations_json=[])
    return ExportService().generate_analysis_pdf_html(analysis, COMPANY)


def _workbook(dataset: dict):
    return load_workbook(io.BytesIO(build_analysis_workbook(deepcopy(dataset))))


def _text(fragment: str) -> str:
    return re.sub(r"\s+", " ", html_lib.unescape(re.sub(r"<[^>]+>", "", fragment))).strip()


def _methods(html: str) -> str:
    return html[html.index('<section class="methods">'):]


def _overview(book) -> list[str]:
    return [str(row[0]) for row in book["Overview"].iter_rows(values_only=True) if row[0]]


def _sources(book) -> list[dict]:
    sheet = book["Sources & methods"]
    header = [cell.value for cell in sheet[1]]
    return [dict(zip(header, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]


def _strip(value):
    """The dataset as an export before the producer sees it: no evidence field at any depth."""
    if isinstance(value, dict):
        return {key: _strip(item) for key, item in value.items() if key not in EVIDENCE_KEYS}
    return [_strip(item) for item in value] if isinstance(value, list) else value


def _expected_line(expected: dict) -> str | None:
    """The line from the fixture's expected places, the page's laterFilingPlaces, in a detail entry's form."""
    def listed(refs: list[dict], suffix: str = "") -> str:
        groups: dict[str, list[str]] = {}
        for ref in refs:
            groups.setdefault(ref["label"] + suffix, []).append(ref["period"])
        return "; ".join(f"{label} ({', '.join(periods)})" for label, periods in groups.items())

    values, growth = expected["values"], expected["comparisons"]
    if not values:
        return None
    count = f"{len(values)} {'figure' if len(values) == 1 else 'figures'}"
    return " ".join(filter(None, [LATER_FILING_MESSAGE, f"{listed(values)}.",
                                  f"{count}; growth across the change reads n/m." if growth else f"{count}.",
                                  f"Comparisons: {listed(growth, ' YoY')}." if growth else ""]))


def _carriers(dataset: dict) -> list[dict]:
    """Every figure and input the exports print that carries its own annual report's value: the window's points and
    the inputs of their value, YoY and window calculations, at any depth."""
    found = []

    def inputs(provenance):
        for operand in (provenance or {}).get("inputs") or []:
            if operand.get("first_reported") and operand.get("value") is not None:
                found.append(operand)
            inputs(operand.get("provenance"))

    keys = {period["key"] for period in dataset["periods"]}
    for series in dataset["series"]:
        for point in series["points"]:
            if point["period"] not in keys:
                continue
            if point.get("first_reported") and point.get("value") is not None:
                found.append(point)
            inputs(point.get("provenance"))
            inputs(point.get("yoy_provenance"))
        inputs(series.get("cagr_provenance"))
        inputs(series.get("window_pp_provenance"))
    return found


def _figure_cells(book, dataset: dict) -> list[tuple]:
    """Every workbook cell that shows a figure, as (sheet, cell, series, period): the Metrics grid's, and the sheets with
    charts', each column read from its header (the top line's by its label, a panel's through _LINE_PANELS)."""
    periods = [period["key"] for period in dataset["periods"]]
    by_label = {series["label"]: series for series in dataset["series"]}
    by_concept = {series["concept"]: series for series in dataset["series"]}
    panels = {title: {label: concept for concept, label in concepts} for title, concepts in _LINE_PANELS}
    cells = [("Metrics", book["Metrics"].cell(row, 4 + offset), series, period)
             for row, series in enumerate(dataset["series"], start=2) for offset, period in enumerate(periods)]
    for sheet in book.worksheets[3:]:
        header = [cell.value for cell in sheet[1]]
        columns = ([by_concept[panels[sheet.title][name]] for name in header[1:]] if sheet.title in panels
                   else [by_label[header[1]]])  # the top line's sheet: Period, its figure, then its YoY growth
        cells += [(sheet.title, sheet.cell(row, column), series, period)
                  for column, series in enumerate(columns, start=2) for row, period in enumerate(periods, start=2)]
    return cells


# ---- 1. The line ----------------------------------------------------------------------------------------------------
@pytest.mark.parametrize("case", CASES, ids=IDS)
def test_both_files_print_the_line_with_the_shared_fixtures_places(case):
    line = _expected_line(case["expected"])
    html = _pdf(case["dataset"])
    entries = re.findall(r"<li(?: id=\"[^\"]*\")?>.*?</li>", re.search(r'<ul class="data-notes">(.*?)</ul>', html, re.S)[1],
                         re.S) if '<ul class="data-notes">' in html else []
    printed = [entry for entry in entries if LATER_FILING_MESSAGE in entry]
    overview = _overview(_workbook(case["dataset"]))
    if line is None:
        assert printed == [] and not any(LATER_FILING_MESSAGE in text for text in overview)
        return
    # The PDF: the detail's last entry, with an empty letter column and no anchor (nothing links to it).
    assert printed == [entries[-1]] == [f'<li><span class="note-letter"></span>{html_lib.escape(line)}</li>']
    # The workbook: one Overview line after the notes' entries, before the closing line.
    closing = workbook_closing(True)
    assert [text for text in overview if LATER_FILING_MESSAGE in text] == [line]
    assert overview.index(line) + 1 == overview.index(closing)


def test_only_the_later_filing_entry_goes_without_an_anchor():
    """A note keeps the anchor main printed for it, whatever its id; only the line, which has no id, prints none."""
    note = {"id": "", "letter": "a", "kind": "review", "reasons": ["source_conflict"], "message": "M.", "values": [],
            "comparisons": []}
    later = later_filing_entry(_case("evidence_with_no_change_of_basis")["dataset"])
    html = ExportService()._analysis_notes_list([note], later)
    assert (html.count('<li id="">'), html.count("<li>")) == (1, 1)


def test_regn_reads_as_the_page_reads_it():
    """WP5b's REGN shape: three figures, and six growths that read n/m across the change (revenue's and the
    calculated net margin's, whose input carries the evidence; never total assets', whose 0.24% gap moves no growth)."""
    overview = _overview(_workbook(_case("regn_own_reports_beside_later_comparatives")["dataset"]))
    assert ("A later filing reports a different value. Revenue (FY2018, FY2019); Total assets (FY2018). 3 figures; "
            "growth across the change reads n/m. Comparisons: Revenue YoY (FY2018, FY2019, FY2020); Net margin YoY "
            "(FY2018, FY2019, FY2020).") in overview


# ---- 2. The evidence ------------------------------------------------------------------------------------------------
REGN_REVENUE_2018 = ("First reported as 6,710,800,000 USD in its own annual report (10-K, filed 2019-02-07). A later "
                     "filing (filed 2021-02-08) reports 5,145,600,000 USD, shown here.")


@pytest.mark.parametrize("case", CASES, ids=IDS)
def test_every_figure_and_input_with_evidence_gives_the_dialogs_sentence_in_both_files(case):
    dataset = case["dataset"]
    carriers = _carriers(dataset)
    methods = _methods(_pdf(dataset))
    book = _workbook(dataset)
    rows = [row for row in _sources(book) if str(row["Entry"]).lower().endswith("first reported")]
    # One evidence row per carrier in each file's Sources & methods, the own report's value as filed beside it.
    assert len(re.findall(r"First reported as ", methods)) == len(rows) == len(carriers), case["name"]
    assert sorted((row["Value (raw)"], row["Unit"], row["Accession"], row["Filed"], row["Source URL"]) for row in rows) == sorted(
        (first["value"], first["unit"], first["accession"], first["filed_at"], first["source_url"])
        for first in (carrier["first_reported"] for carrier in carriers))
    for row in rows:
        assert (row["Method"], row["Note"].startswith("First reported as ")) == ("reported", True), row
    # Each cell that shows a figure carries the sentence exactly where the figure carries the evidence: in Metrics, on
    # the top line's sheet and on each line panel's.
    shown = _figure_cells(book, dataset)
    assert {sheet for sheet, *_ in shown} == {"Metrics", *(sheet.title for sheet in book.worksheets[3:])}
    for sheet, cell, series, period in shown:
        point = next((point for point in series["points"] if point["period"] == period), {})
        own = (point.get("first_reported") or {}).get("value")
        words = f"First reported as {own:,.0f} {series['unit']}" if own is not None else "First reported as "
        assert (cell.comment is not None and words in cell.comment.text) is (own is not None), (sheet, cell.coordinate)


def test_the_sentence_names_both_filings_and_the_value_shown():
    dataset = _case("regn_own_reports_beside_later_comparatives")["dataset"]
    book = _workbook(dataset)
    assert book["Metrics"]["E2"].comment.text == REGN_REVENUE_2018
    assert book["Revenue & growth"]["B3"].comment.text == REGN_REVENUE_2018
    first = next(row for row in _sources(book) if (row["Metric"], row["Period"], row["Entry"]) == ("Revenue", "FY2018", "First reported"))
    assert (first["Value (raw)"], first["Note"]) == (6_710_800_000, REGN_REVENUE_2018)
    assert first["Source URL"] == "https://www.sec.gov/Archives/edgar/data/872589/000153217619000009/"
    # A calculated figure's input carries it as the figure does (item 5): net margin FY2018's revenue, in its rows.
    margin = [row["Entry"] for row in _sources(book) if (row["Metric"], row["Period"]) == ("Net margin", "FY2018")]
    assert margin[:4] == ["Value", "Input 1", "Input 2", "Input 2 first reported"]
    methods = _methods(_pdf(dataset))
    assert (f'<ul class="operands"><li>{html_lib.escape(REGN_REVENUE_2018)} · Accession: 0001532176-19-000009 · '
            '<a href="https://www.sec.gov/Archives/edgar/data/872589/000153217619000009/">Source filing 1</a></li></ul>'
            ) in methods
    # Under the figure's own row in the PDF, before its YoY's.
    revenue = methods[methods.index("<h3>Revenue</h3>"):methods.index("<h3>Net income</h3>")]
    rows = re.findall(r"<li><strong>(FY\d{4} · \w+)</strong>(.*?)(?=<li><strong>|$)", revenue, re.S)
    assert [label for label, body in rows if "First reported as" in body] == ["FY2018 · Value", "FY2019 · Value"]


def test_a_finding_says_what_was_found():
    dataset = _case("a_sign_flip_and_a_scale_factor_count_as_figures")["dataset"]
    book = _workbook(dataset)
    sign, scale = book["Metrics"]["E2"].comment.text.split("\n"), book["Metrics"]["D3"].comment.text.split("\n")
    # The note's letter first (the mark is the producer's source_conflict, D2), then the evidence and the finding.
    assert sign[0].startswith("Data note a: ") and scale[0].startswith("Data note a: ")
    assert sign[1:] == ["First reported as 1,895,000,000 USD in its own annual report (10-K, filed 2021-02-08). A later "
                        "filing (filed 2022-02-07) reports -1,895,000,000 USD, shown here.",
                        "A later filing reports the opposite sign."]
    assert scale[2] == "A later filing reports a value 1,000 times larger than the year’s own annual report."
    assert "A later filing reports the opposite sign." in _methods(_pdf(dataset))
    # The factor's size and direction come from the two values.
    debt = _point(dataset, "long_term_debt", "FY2019")
    for own, words in ((5_969_900_000_000.0, "1,000 times smaller"), (5_969.9, "1,000,000 times larger")):
        debt["first_reported"]["value"] = own
        assert f"A later filing reports a value {words} than the year’s own annual report." in _workbook(dataset)["Metrics"]["D3"].comment.text


def _page(path: str, pattern: str, what: str) -> re.Match:
    """One of the page's evidence templates, read from its source; a change of its code form fails until re-read."""
    found = re.search(pattern, (ANALYSIS / path).read_text(encoding="utf-8"))
    assert found, f"{path}'s {what} moved: re-read it and the files' copy in analysis_export_metadata.py"
    return found


def test_the_evidence_words_are_the_pages():
    """The words the files print of a figure's evidence are the page dialog's (D11, D12), read from its sources:
    laterFilings.ts firstReportedSentence and filingParts, disagreementSentence's two findings, and SourceValue.tsx's
    cause and "As first reported:". Only the values differ, as chief design approved on #1866: exact figures for the
    page's compact ones, and days as stored for "7 Feb 2019". A reworded page fails here until the files follow."""
    sentence = _page("lib/laterFilings.ts", r"return `(First reported as )\$\{own\}( in its own annual report)\$\{ownFiling && `"
                     r"( \()\$\{ownFiling\}(\))`\}(\. A later filing)\$\{laterFiled && `( \(filed )\$\{laterFiled\}(\))`\}"
                     r"( reports )\$\{later\}(, shown here\.)`", "firstReportedSentence")
    parts = _page("lib/laterFilings.ts", r"\[form, filed && `(filed )\$\{filed\}`\]\.filter\(Boolean\)\.join\('([^']*)'\)",
                  "filingParts")
    sign = _page("lib/laterFilings.ts", r"disagreement === 'opposite_sign'\) return '([^']*)'", "sign finding")
    scale = _page("lib/laterFilings.ts", r"return `(A later filing reports a value )\$\{fmtCount\(1000 \*\* Math\.abs\(steps\)\)\}"
                  r"( times )\$\{steps > 0 \? '([^']*)' : '([^']*)'\}( than [^`]*)`", "scale finding")
    cause = _page("components/SourceValue.tsx", r"\{comparison\.laterFiling && <p>([^<]*)</p>\}", "basis cause")
    growth = _page("components/SourceValue.tsx", r"\{comparison\.laterFiling\?\.asFirstReported && <p[^>]*>([^<]*)<GrowthText "
                   r"text=\{comparison\.laterFiling\.asFirstReported\} /></p>\}", "growth as first reported")

    def page_sentence(own: str, filing: list[str], later_filed: str | None, later: str) -> str:
        own_filing = parts[2].join(filing)
        return "".join([sentence[1], own, sentence[2], f"{sentence[3]}{own_filing}{sentence[4]}" if own_filing else "",
                        sentence[5], f"{sentence[6]}{later_filed}{sentence[7]}" if later_filed else "", sentence[8], later,
                        sentence[9]])

    first = {"value": 6_710_800_000.0, "unit": "USD", "form": "10-K", "filed_at": "2019-02-07", "later_filed_at": "2021-02-08"}
    own, later = "6,710,800,000 USD", "5,145,600,000 USD"
    assert first_reported_sentence(first, 5_145_600_000.0) == page_sentence(own, ["10-K", f"{parts[1]}2019-02-07"],
                                                                             "2021-02-08", later)
    bare = {**first, "form": None, "filed_at": None, "later_filed_at": None}
    assert first_reported_sentence(bare, 5_145_600_000.0) == page_sentence(own, [], None, later)
    assert disagreement_sentence({**first, "disagreement": "opposite_sign"}, -6_710_800_000.0) == sign[1]
    for shown, factor, direction in ((6_710_800_000_000.0, "1,000", scale[3]), (6_710.8, "1,000,000", scale[4])):
        assert disagreement_sentence({**first, "disagreement": "scale"}, shown) == (
            f"{scale[1]}{factor}{scale[2]}{direction}{scale[5]}")
    assert evidence_lines({"first_reported_growth": 0.1428084874493376, "percent": False}) == [
        cause[1], f"{growth[1]}{growth_text(0.1428084874493376, False)}"]
    assert BASIS_CHANGE_CAUSE == cause[1]


GROWTH = {
    "regn_own_reports_beside_later_comparatives": {
        ("revenue", "FY2018"): "+14.3%", ("revenue", "FY2019"): "+17.2%", ("revenue", "FY2020"): "+8.1%",
        ("net_margin", "FY2018"): "+16.0pp", ("net_margin", "FY2019"): "-9.5pp", ("net_margin", "FY2020"): "+14.4pp"},
    # Across a sign or scale finding the page prints n/m (chief design on #1776).
    "a_sign_flip_and_a_scale_factor_count_as_figures": {
        ("financing_cash_flow", "FY2020"): "n/m", ("long_term_debt", "FY2020"): "n/m"},
    # Null: the own reports give no growth, so only the cause prints.
    "a_null_growth_as_first_reported_counts": {("net_income", "FY2018"): "-100.0%", ("net_income", "FY2019"): None},
}


@pytest.mark.parametrize("name", list(GROWTH))
def test_growth_across_the_change_gives_its_cause_and_growth_as_first_reported(name):
    dataset = _case(name)["dataset"]
    labels = {series["concept"]: series["label"] for series in dataset["series"]}
    rows = {(row["Metric"], row["Period"]): row["Note"] for row in _sources(_workbook(dataset)) if row["Entry"] == "YOY"}
    methods = _methods(_pdf(dataset))
    for (concept, period), growth in GROWTH[name].items():
        words = BASIS_CHANGE_CAUSE + (f" As first reported: {growth}" if growth else "")
        note = rows[(labels[concept], period)]
        assert note.split(" · ")[-1] == words, (concept, period, note)
        assert f"<strong>{period} · YOY</strong>" in methods and f'<ul class="operands"><li>{words}</li></ul>' in methods
    # Every YoY the case does not list says nothing of a later filing: a sign change's n/m among them.
    assert sum(BASIS_CHANGE_CAUSE in note for note in rows.values() if note) == len(GROWTH[name])
    assert methods.count(BASIS_CHANGE_CAUSE) == len(GROWTH[name])


def test_the_top_lines_growth_comment_gives_the_cause():
    book = _workbook(_case("regn_own_reports_beside_later_comparatives")["dataset"])
    top = book["Revenue & growth"]
    assert (top["C3"].value, top["C3"].number_format) == ("n/m", "General")
    assert top["C3"].comment.text.split("\n")[:2] == [BASIS_CHANGE_CAUSE, "As first reported: +14.3%"]
    assert top["C5"].comment.text.split("\n")[1] == "As first reported: +8.1%"
    # FY2021 keeps its growth and says nothing of a later filing.
    assert isinstance(top["C6"].value, float) and BASIS_CHANGE_CAUSE not in top["C6"].comment.text


GROWTH_TABLE = json.loads((Path(__file__).resolve().parents[1] / "fixtures" / "analysis_growth_text.json").read_text())
GROWTH_TEXT = [(group["rule"], value, percent, text) for group in GROWTH_TABLE["groups"]
               for value, percent, text in group["cases"]]


@pytest.mark.parametrize(("rule", "value", "percent", "text"), GROWTH_TEXT,
                         ids=[f"{value}-{'pp' if percent else 'fraction'}" for _, value, percent, _ in GROWTH_TEXT])
def test_the_growth_as_first_reported_reads_as_the_page_prints_it(rule, value, percent, text):
    """The page's formatGrowth (growth.ts) over the table both copies run; frontend/tests/unit/analysis-growth-text.spec.ts
    runs the page's own copy on it, so the table cannot drift from the page."""
    assert growth_text(value, percent) == text, rule


def test_the_files_print_the_growth_through_the_pages_rule():
    """Zero, a growth past 1,000% and a tie reach both files as the page prints them, never as the prompt's formatter
    prints them ("+0.0%", "+1234.5%", "+0.2%")."""
    dataset = _case("regn_own_reports_beside_later_comparatives")["dataset"]
    for period, growth in (("FY2018", 0.0), ("FY2019", 12.345), ("FY2020", 0.0025)):
        _point(dataset, "revenue", period)["yoy_first_reported"] = growth
    top = _workbook(dataset)["Revenue & growth"]
    printed = [top[cell].comment.text.split("\n")[1] for cell in ("C3", "C4", "C5")]
    assert printed == ["As first reported: 0.0%", "As first reported: +1,234.5%", "As first reported: +0.3%"]
    methods = _methods(_pdf(dataset))
    assert all(f"{BASIS_CHANGE_CAUSE} {words}</li>" in methods for words in printed)


def test_a_growth_is_n_m_across_the_change_only_where_it_reads_n_m():
    """laterFilings.ts basisChanged: the key beside a YoY that is not n/m says nothing, in either file."""
    dataset = _case("evidence_with_no_change_of_basis")["dataset"]
    _point(dataset, "total_assets", "FY2019")["yoy_first_reported"] = 0.12
    assert BASIS_CHANGE_CAUSE not in _pdf(dataset)
    assert not any(BASIS_CHANGE_CAUSE in str(row["Note"]) for row in _sources(_workbook(dataset)))


def test_a_figure_with_no_value_gives_no_evidence():
    """The page's dialog reads first_reported only beside a value (SourceValue.tsx), and so do the files."""
    dataset = _case("evidence_with_no_change_of_basis")["dataset"]
    _point(dataset, "total_assets", "FY2018")["value"] = None
    assert "First reported as" not in _pdf(dataset)
    book = _workbook(dataset)
    assert not any("first reported" in str(row["Entry"]).lower() for row in _sources(book))
    assert "First reported" not in book["Metrics"]["E2"].comment.text  # the gap's own note stays


# ---- 3. The sentences -----------------------------------------------------------------------------------------------
@pytest.mark.parametrize("case", CASES, ids=IDS)
def test_the_closing_and_n_m_sentences_print_where_the_line_and_the_n_m_do(case):
    later, growth = bool(case["expected"]["values"]), bool(case["expected"]["comparisons"])
    html = _pdf(case["dataset"])
    overview = _overview(_workbook(case["dataset"]))
    assert f'<p class="fine-print">{html_lib.escape(pdf_closing(later))}</p>' in html
    assert (LATER_FILING_EVIDENCE in html_lib.unescape(html)) is later
    assert overview.count(workbook_closing(later)) == 1
    assert (WORKBOOK_CLOSING in overview) is not later
    # Sources & methods' n/m sentence closes its footnote in the PDF and follows its lines in the Overview.
    footnote = _text(re.search(r'<section class="methods"><h2>Sources &amp; methods</h2>\s*<p class="footnote">(.*?)</p>',
                               html, re.S)[1])
    assert footnote.endswith(NOT_MEANINGFUL_CAUSES) is growth
    assert (NOT_MEANINGFUL_CAUSES in overview) is growth
    if growth:
        assert overview[overview.index(NOT_MEANINGFUL_CAUSES) - 1].startswith("Sources & methods preserves raw")


def test_the_n_m_sentence_names_only_the_causes_the_data_carries():
    """The dataset reads a growth n/m where its two years' values have opposite signs (trend_analysis/dataset.py
    _growth): a zero prior year gives no growth, which the files print as unavailable, and a zero this year gives
    -100%, so the sentence names no zero (D12's draft, "where it changes sign or starts from zero, or where …", named
    one, and its "it" read as the growth's own sign); a later filing's value is the other cause."""
    from app.services.trend_analysis.dataset import _growth

    assert (_growth(5.0, 0.0), _growth(-1.0, 1.0), _growth(1.0, -1.0), _growth(0.0, 5.0)) == (None, "nm", "nm", -1.0)
    assert NOT_MEANINGFUL_CAUSES == ("Growth reads n/m where its two years’ values have opposite signs, or where a later "
                                     "filing reports a different value for one or both of its years.")


def test_the_closing_lines_add_only_the_evidence_sentence_and_the_comment_clause():
    assert pdf_closing(False) == PDF_CLOSING and workbook_closing(False) == WORKBOOK_CLOSING
    assert pdf_closing(True) == PDF_CLOSING.replace(" Automated", f" {LATER_FILING_EVIDENCE} Automated", 1)
    assert workbook_closing(True) == WORKBOOK_CLOSING.replace(" In Metrics", f" {LATER_FILING_EVIDENCE} In Metrics", 1).replace(
        "a data note or a calculation", "a data note, a calculation or a first-reported value")


# ---- 4. No letter ---------------------------------------------------------------------------------------------------
@pytest.mark.parametrize("case", CASES, ids=IDS)
def test_evidence_marks_nothing(case):
    letters = {note["letter"] for note in data_notes(case["dataset"])["notes"] if marks_figure(note)}
    assert pdf_letters(_pdf(case["dataset"])) == letters
    assert workbook_letters(build_analysis_workbook(deepcopy(case["dataset"]))) == letters


# ---- 5. Parts left out ----------------------------------------------------------------------------------------------
@pytest.mark.parametrize(("missing", "sentence"), [
    (("form",), "First reported as 6,710,800,000 USD in its own annual report (filed 2019-02-07). A later filing (filed "
                "2021-02-08) reports 5,145,600,000 USD, shown here."),
    (("filed_at",), "First reported as 6,710,800,000 USD in its own annual report (10-K). A later filing (filed 2021-02-08) "
                    "reports 5,145,600,000 USD, shown here."),
    (("form", "filed_at", "later_filed_at"), "First reported as 6,710,800,000 USD in its own annual report. A later filing "
                                             "reports 5,145,600,000 USD, shown here."),
], ids=["no-form", "no-filed-day", "no-form-or-days"])
def test_a_part_the_payload_lacks_is_left_out(missing, sentence):
    dataset = _case("regn_own_reports_beside_later_comparatives")["dataset"]
    first = _point(dataset, "revenue", "FY2018")["first_reported"]
    first.update({part: None for part in missing})
    assert _workbook(dataset)["Metrics"]["E2"].comment.text == sentence
    assert f"<li>{html_lib.escape(sentence)} · Accession: " in _pdf(dataset)


def test_a_missing_link_is_left_out():
    dataset = _case("regn_own_reports_beside_later_comparatives")["dataset"]
    _point(dataset, "revenue", "FY2018")["first_reported"]["source_url"] = None
    assert f"<li>{html_lib.escape(REGN_REVENUE_2018)} · Accession: 0001532176-19-000009</li>" in _pdf(dataset)
    row = next(row for row in _sources(_workbook(dataset)) if (row["Period"], row["Entry"]) == ("FY2018", "First reported"))
    assert row["Source URL"] in (None, "")


def test_a_per_share_figure_reads_per_share():
    dataset = _case("evidence_with_no_change_of_basis")["dataset"]
    assets = next(series for series in dataset["series"] if series["concept"] == "total_assets")
    assets.update(unit="USD/shares")
    point = _point(dataset, "total_assets", "FY2018")
    point.update(value=2.05, unit="USD/shares")
    point["first_reported"].update(value=2.1, unit="USD/shares")
    assert _workbook(dataset)["Metrics"]["E2"].comment.text.startswith(
        "First reported as 2.1 USD per share in its own annual report (10-K, filed 2019-02-07). A later filing (filed "
        "2020-02-07) reports 2.05 USD per share, shown here.")


# ---- 6. Inert until the producer ------------------------------------------------------------------------------------
NEW_WORDS = (LATER_FILING_MESSAGE, "First reported", "first reported", BASIS_CHANGE_CAUSE, NOT_MEANINGFUL_CAUSES,
             LATER_FILING_EVIDENCE, "first-reported value")


@pytest.mark.parametrize("case", CASES, ids=IDS)
def test_a_dataset_without_the_fields_prints_none_of_it(case):
    dataset = _strip(case["dataset"])
    html = html_lib.unescape(_pdf(dataset))
    book = _workbook(dataset)
    texts = [str(cell.value) for sheet in book.worksheets for row in sheet.iter_rows() for cell in row if cell.value]
    texts += [cell.comment.text for sheet in book.worksheets for row in sheet.iter_rows() for cell in row if cell.comment]
    assert [words for words in NEW_WORDS if words in html or any(words in text for text in texts)] == []
    assert f'<p class="fine-print">{PDF_CLOSING}</p>' in html and WORKBOOK_CLOSING in texts
