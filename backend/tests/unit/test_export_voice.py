"""The exports' rendered words keep the voice (docs/voice-and-style.md; CLAUDE.md rule 12).

The frontend's copy gates (no-em-dash-copy.spec.ts, brandCopyRules.spec.ts) read TypeScript only, and a
walk over Python literals cannot see what a PDF, a CSV or a workbook prints once the fixtures have
flowed through summary_sections and the export builders. This test renders the filing-summary PDF and
CSV, the analysis PDF and the workbook (its cells, comments, sheet names and chart titles) for fixtures
that reach every fixed line the four export modules write, and fails on:

1. an em-dash, unless the cell or text run is the bare null token "—" (the one sanctioned em-dash);
2. an exclamation mark or an emoji (a pictograph, a dingbat or the emoji variation selector);
3. a window key's ".." where a range prints ("FY2016..FY2025", the stored form the CAGR sentence
   carries): the voice guide's range takes the en-dash, "FY2016–FY2025";
4. a heading the export modules own that is not in sentence case: the PDFs' titles, the analysis PDF's
   section headings (its narrative carries none here, so every <h2> is the builder's), the CSV's title
   row, and the workbook's title, sheet names, header rows and fact labels. A proper noun passes (the
   company's name, "Multi-Period Analysis", "EarningsNerd", "YoY") and so does any word with no
   lowercase letter (SEC, UTC, 10-K, FY2024). The Metrics header row ends in the key to the letters the grid
   prints (D-061), one cell past the grid with no header fill, which names metrics by their labels: there, and
   only there, the run's metric labels stand as proper nouns;
5. EDGAR, in any case and inside a word too ("SEC_EDGAR", "EDGARs"), anywhere they print outside an
   address: the exports name their source SEC.gov, as the sign-in page and the filing page do ("Original on
   SEC.gov · {url}", the filing page's words for its link to the primary document, with the address that
   link opens; E190), since no acronym stands in a label (voice rule 2). An address (http or https, up to
   the first character outside printable ASCII, such as a space or the middot) is not copy, so a filing's
   /edgar/ path passes, and the fixtures quote nothing that names EDGAR, so every text this reads is the
   exports' own.

Out of reach (review): the summary's section titles and table headings, which summary_sections.py
writes in the schema's casing (the voice guide keeps it); and lines that only other data reaches.
"""
from __future__ import annotations

import csv
import html as html_lib
import io
import re
import zipfile
from html.parser import HTMLParser

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

import tests.unit.test_excel_export as workbook_fixtures
import tests.unit.test_export_service as export_fixtures
from app.services.excel_export_service import build_analysis_workbook
from app.services.export_service import ExportService
from app.services.trend_analysis.data_notes import data_notes

EM_DASH = "—"
EDGAR = re.compile(r"edgar", re.IGNORECASE)
ADDRESS = re.compile(r"https?://[!-~]+")
WINDOW_KEY = re.compile(r"(?:FY)?\d{4}(?:Q[1-4])?\.\.(?:FY)?\d{4}|Q[1-4] \d{4}\.\.Q[1-4] \d{4}")
EMOJI = re.compile("[\u2600-\u27bf\u2b00-\u2bff\ufe0f\U0001f000-\U0001faff]")
PROPER = ("BIOGEN INC.", "Test & Co", "Multi-Period Analysis", "EarningsNerd")
MIXED_CASE = frozenset({"YoY", "QoQ"})
_WORD = re.compile(r"[A-Za-z][A-Za-z0-9'’-]*")
_VOID = frozenset({"br", "meta", "img", "hr", "link", "input", "wbr"})


class _Runs(HTMLParser):
    """A document's text runs (the stylesheet's content strings among them) and its h1/h2 texts."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.runs: list[str] = []
        self.headings: list[str] = []
        self._open: list[str] = []
        self._heading: list[str] | None = None

    def handle_starttag(self, tag, attrs):
        if tag in ("h1", "h2"):
            self._heading = []
        if tag not in _VOID:
            self._open.append(tag)

    def handle_endtag(self, tag):
        if tag in ("h1", "h2") and self._heading is not None:
            self.headings.append("".join(self._heading).strip())
            self._heading = None
        while self._open and self._open.pop() != tag:
            pass

    def handle_data(self, data):
        if self._open and self._open[-1] == "style":
            # Margin-box copy, such as the running footer, prints from the stylesheet's content strings.
            self.runs += re.findall(r'"([^"]*)"', " ".join(re.findall(r"content:\s*([^;]+);", data)))
            return
        if data.strip():
            self.runs.append(data)
        if self._heading is not None:
            self._heading.append(data)


def read_html(document: str) -> _Runs:
    parser = _Runs()
    parser.feed(document)
    return parser


def breaks(texts: list[str]) -> list[str]:
    """Texts that carry an em-dash other than the bare null token, an exclamation mark, an emoji or a
    window key where a range prints."""
    return [text for text in texts
            if (EM_DASH in text and text.strip() != EM_DASH) or "!" in text or EMOJI.search(text)
            or WINDOW_KEY.search(text)]


def names_edgar(texts: list[str]) -> list[str]:
    """Texts that name EDGAR in any case outside an address (check 5)."""
    return [text for text in texts if EDGAR.search(ADDRESS.sub(" ", text))]


def not_sentence_case(text: str) -> str | None:
    """The first word that breaks sentence case, or None. A proper noun stands in as one capital word."""
    for name in PROPER:
        text = text.replace(name, " NAME ")
    for index, word in enumerate(_WORD.findall(text)):
        if word in MIXED_CASE or not any(c.islower() for c in word):
            continue
        if word[0].isupper() != (index == 0):
            return word
    return None


def _summary():
    sections = export_fixtures._full_sections()
    sections["management_discussion_insights"]["quotes"].append({"quote": "Demand held up in every region."})
    return export_fixtures._make_summary_and_filing(sections)


def _analysis():
    """The PDF fixture with a narrative that carries no heading, a calculated point with its operand's
    source fields, a metric calculated in full (the sentence above the grid names it), and data notes in every
    count the summary line prints."""
    analysis = export_fixtures.TestAnalysisPdfHtml()._analysis()
    # The CAGR sentence as observations.py stores it, with the window key.
    analysis.narrative_md = "Revenue grew from $1,000 [1] to $1,500 [2]. Revenue CAGR was 22.5% over FY2022..FY2024 [3]."
    points = analysis.dataset_json["series"][0]["points"]
    points[0]["provenance"] = {"method": "reported", "validation": "needs_review", "reasons": ["source_conflict"]}
    points[1].update(yoy=0.2, yoy_provenance={"method": "calculated", "validation": "needs_review",
                                              "reasons": ["source_conflict"]})
    points[2]["provenance"] = {
        "method": "calculated", "validation": "passed", "reasons": [], "formula": "annual_minus_ytd",
        "inputs": [{"concept": "operating_cash_flow", "value": 2500, "unit": "USD", "period_start": "2024-01-01",
                    "period_end": "2024-12-31", "accession": "0000123456-25-000001", "raw_tag": "us-gaap:Revenues",
                    "filed_at": "2025-02-14", "source_url": "https://www.sec.gov/Archives/edgar/data/1/000012345625000001/"}],
    }
    analysis.dataset_json["series"].append({
        "concept": "free_cash_flow", "label": "Free cash flow", "unit": "USD", "percent": False,
        "points": [{"period": period, "value": 200.0, "provenance": {"method": "calculated", "validation": "passed"}}
                   for period in ("FY2022", "FY2023", "FY2024")],
    })
    analysis.dataset_json["data_notes"] = data_notes(analysis.dataset_json)
    return analysis


def _quarterly_analysis():
    """A quarterly run whose Q4 flow figure is calculated: its grid's Q4 header takes the † and the legend explains it."""
    analysis = export_fixtures.TestAnalysisPdfHtml()._analysis()
    analysis.mode, analysis.period_key = "quarterly", "2024Q3..2024Q4"
    analysis.dataset_json.update(mode="quarterly", period_key="2024Q3..2024Q4", periods=[
        {"key": "2024Q3", "fiscal_period": "Q3"}, {"key": "2024Q4", "fiscal_period": "Q4"}], series=[
        {"concept": "revenue", "label": "Revenue", "unit": "USD", "points": [
            {"period": "2024Q3", "value": 300.0}, {"period": "2024Q4", "value": 400.0, "derived": True}]}])
    analysis.dataset_json["data_notes"] = data_notes(analysis.dataset_json)
    return analysis


def _dataset():
    """The workbook fixture with a note that questions a figure, so the Metrics header row prints its key."""
    dataset = workbook_fixtures._dataset()
    dataset["series"][0]["points"][0]["provenance"] = {"method": "reported", "validation": "needs_review",
                                                       "reasons": ["source_conflict"]}
    return dataset


def _workbook():
    return load_workbook(io.BytesIO(build_analysis_workbook(_dataset(), exported_at=workbook_fixtures.EXPORTED_AT)))


def metrics_header_and_key(sheet, labels: list[str]) -> tuple[list[str], str | None]:
    """The Metrics header row's cells, and its key with each metric label read as a proper noun. The key is exactly the
    cell one column past the last header, without the header fill; every other cell of the row reads as a header. None
    when the grid prints no letter."""
    row = [cell for cell in sheet[1] if cell.value]
    past = max((cell.column for cell in row if cell.fill.fill_type is not None), default=0) + 1
    key_cell = next((cell for cell in row if cell.column == past and cell.fill.fill_type is None), None)
    if key_cell is None:
        return [cell.value for cell in row], None
    key = key_cell.value
    for label in sorted(labels, key=len, reverse=True):
        key = key.replace(label, " NAME ")
    return [cell.value for cell in row if cell is not key_cell], key


def _workbook_texts(data: bytes) -> list[str]:
    book = load_workbook(io.BytesIO(data))
    texts = list(book.sheetnames)
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    texts.append(cell.value)
                if cell.comment is not None:
                    texts.append(cell.comment.text)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        for name in archive.namelist():
            if name.startswith("xl/charts/"):
                texts += re.findall(r"<a:t>([^<]*)</a:t>", archive.read(name).decode("utf-8"))
    return texts


def test_no_em_dash_exclamation_emoji_or_window_key_in_what_the_exports_print():
    service = ExportService()
    summary, filing = _summary()
    analysis = _analysis()
    found = {
        "filing PDF": breaks(read_html(service.generate_pdf_html(summary, filing)).runs),
        "analysis PDF": breaks(read_html(service.generate_analysis_pdf_html(
            analysis, export_fixtures.TestAnalysisPdfHtml()._company())).runs),
        "quarterly analysis PDF": breaks(read_html(service.generate_analysis_pdf_html(
            _quarterly_analysis(), export_fixtures.TestAnalysisPdfHtml()._company())).runs),
        "CSV": breaks([cell for row in csv.reader(io.StringIO(service.generate_csv(summary, filing))) for cell in row]),
        "workbook": breaks(_workbook_texts(build_analysis_workbook(_dataset(), exported_at=workbook_fixtures.EXPORTED_AT))),
    }
    assert found == {name: [] for name in found}


def test_the_exports_name_their_source_sec_gov_not_edgar():
    service = ExportService()
    summary, filing = _summary()
    found = {
        "filing PDF": names_edgar(read_html(service.generate_pdf_html(summary, filing)).runs),
        "analysis PDF": names_edgar(read_html(service.generate_analysis_pdf_html(
            _analysis(), export_fixtures.TestAnalysisPdfHtml()._company())).runs),
        "CSV": names_edgar([cell for row in csv.reader(io.StringIO(service.generate_csv(summary, filing))) for cell in row]),
        "workbook": names_edgar(_workbook_texts(build_analysis_workbook(_dataset(), exported_at=workbook_fixtures.EXPORTED_AT))),
    }
    assert found == {name: [] for name in found}
    url = "https://www.sec.gov/Archives/edgar/data/875045/x.htm"
    named = [f"Source: SEC EDGAR · {url}", "SEC-EDGAR", f"Source: SEC Edgar · {url}", "Open on edgar",
             f"Source: SEC_EDGAR · {url}", "EDGARs", f"{url} · SEC EDGAR", f"{url}·EDGAR"]
    assert names_edgar([*named, f"Original on SEC.gov · {url}", url]) == named


def test_the_headings_the_exports_own_are_in_sentence_case():
    service = ExportService()
    summary, filing = _summary()
    filing_pdf = read_html(service.generate_pdf_html(summary, filing))
    analysis_pdf = read_html(service.generate_analysis_pdf_html(_analysis(), export_fixtures.TestAnalysisPdfHtml()._company()))
    csv_rows = list(csv.reader(io.StringIO(service.generate_csv(summary, filing).removeprefix("\ufeff"))))
    book = _workbook()
    overview = book["Overview"]
    metrics_header, key = metrics_header_and_key(book["Metrics"], [series["label"] for series in _dataset()["series"]])
    assert key is not None, "the fixture's grid prints a letter, so its header row ends in the key"
    headings = {
        "filing PDF title": [filing_pdf.headings[0]],
        "analysis PDF headings": analysis_pdf.headings,
        "CSV title row": csv_rows[3],
        "workbook sheet names": book.sheetnames,
        "workbook title": [overview["A7"].value],
        "workbook fact labels": [row[0].value for row in overview.iter_rows(min_row=10) if row[1].value],
        "workbook header rows": [*metrics_header, *(cell.value for sheet in book.worksheets[1:] if sheet.title != "Metrics"
                                                    for cell in sheet[1] if cell.value)],
        "workbook Metrics key": [key],
    }
    # The data notes are fine print under the grid, with no heading of their own (D-061).
    assert headings["analysis PDF headings"][1:] == [
        "Metrics by period", "Narrative sources", "Sources & methods", "About this document",
    ], "the analysis fixture's narrative must carry no heading, so every <h2> read here is the builder's"
    found = {name: [(text, word) for text in texts if (word := not_sentence_case(html_lib.unescape(str(text))))]
             for name, texts in headings.items()}
    assert found == {name: [] for name in headings}


def test_the_readers_flag_the_forms_this_pr_fixed():
    """Planted forms: the old copy fails each check, and the null token and a sentence-case title pass."""
    assert breaks(["BIOGEN INC. — 10-K Summary", "“We remain focused.” — CEO", "Ready!", "Exported \U0001F389",
                   "Done \u2705", EM_DASH, f" {EM_DASH} ", "1 calculated †", "FY2022–FY2024 · 3 periods"]) == [
        "BIOGEN INC. — 10-K Summary", "“We remain focused.” — CEO", "Ready!", "Exported \U0001F389", "Done \u2705",
    ]
    assert breaks(["Revenue CAGR was 22.5% over FY2016..FY2025 [3].", "2025Q1..2026Q2", "Q1 2025..Q2 2026",
                   "FY2016–FY2025", "Q1 2025–Q2 2026", "Loading...", "10..20"]) == [
        "Revenue CAGR was 22.5% over FY2016..FY2025 [3].", "2025Q1..2026Q2", "Q1 2025..Q2 2026",
    ]
    assert not_sentence_case("BIOGEN INC. · 10-K Summary") == "Summary"
    assert not_sentence_case("Sources & Methods") == "Methods"
    assert not_sentence_case("BIOGEN INC. · 10-K summary") is None
    assert not_sentence_case("Test & Co · Multi-Period Analysis") is None
    assert not_sentence_case("YoY growth") is None and not_sentence_case("Exported UTC") is None
    # The Metrics key: the run's labels stand as proper nouns there; its own words keep sentence case. A header cell
    # is never read as the key, so its capitals still fail.
    sheet = Workbook().active
    for column, text in enumerate(["Metric", "Window Growth"], start=1):
        sheet.cell(1, column, text).fill = PatternFill("solid", start_color="F4F3EE")
    assert metrics_header_and_key(sheet, ["Revenue"]) == (["Metric", "Window Growth"], None)
    sheet.cell(1, 3, "Data notes · a Free cash flow FY2024: not yet matched to the filing")
    header, key = metrics_header_and_key(sheet, ["Free cash flow"])
    assert not_sentence_case(header[1]) == "Growth" and not_sentence_case(key) is None
    # Only that cell: an unfilled cell further along reads as a header, its capitals failing.
    sheet.cell(1, 5, "Data Notes · a Free cash flow")
    assert metrics_header_and_key(sheet, ["Free cash flow"]) == (["Metric", "Window Growth", "Data Notes · a Free cash flow"], key)
    sheet.cell(1, 5).value = None
    # And with no cell one past the headers, an unfilled cell after a gap is no key: it reads as a header.
    gapped = Workbook().active
    for column, text in enumerate(["Metric", "Window Growth"], start=1):
        gapped.cell(1, column, text).fill = PatternFill("solid", start_color="F4F3EE")
    gapped.cell(1, 4, "Data Notes · a Free cash flow")
    assert metrics_header_and_key(gapped, ["Free cash flow"]) == (
        ["Metric", "Window Growth", "Data Notes · a Free cash flow"], None)
    for broken, word in (("Data Notes · a Free cash flow FY2024: not yet matched", "Notes"),
                         ("Data notes · a Free cash flow FY2024: Not yet matched", "Not")):
        sheet.cell(1, 3, broken)
        assert not_sentence_case(metrics_header_and_key(sheet, ["Free cash flow"])[1]) == word
    runs = read_html("<style>@page { @bottom-left { content: \"A — B\"; } }</style><h1>X · Y summary</h1><td>—</td>")
    assert runs.runs == ["A — B", "X · Y summary", EM_DASH] and runs.headings == ["X · Y summary"]
