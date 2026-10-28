"""Every export colour is a design token in its role, and every export text clears 4.5:1 on its ground.

The frontend's colour gates (designSnapshotParity, designTokenParity, the ESLint hex rule) never read
Python, and the exports had drifted: pdf_branding.PALETTE held #6B7280 as the tertiary ink, the
flat-delta grey, which reads 4.26:1 on the masthead's tint, where the tertiary ink token is #636A77.
This test reads DESIGN.md's frontmatter (designSnapshotParity holds it equal to tailwind.config.js)
and checks the exports as they render. A PDF's stylesheet is read as WeasyPrint reads it: every
declaration of every block in order, none collapsed, a rule nested in another (CSS nesting) under
its full selector, a rule inside @media or @supports, @font-face, and @page's margin boxes as
"@page @top-left".

1. Roles: each PALETTE key names one token in ROLES and holds that token's value, so a value that is
   some other token (the flat grey for the tertiary ink) fails; CHART_SERIES is chart-1 to chart-6 in
   order. A new PALETTE key fails until ROLES names its token.
2. Output: every colour value in the two PDFs is a token's hex, `inherit` or `currentColor`, in every
   declaration of the stylesheet and in the SVG's paint attributes. A colour function (rgb(), hsl(),
   color-mix() and the rest), a named, system or `transparent` keyword, a hex with an alpha digit, an
   indirection the gate cannot read (var(), env(), attr()) and a CSS-wide keyword other than inherit
   on a property that carries a colour (initial, unset, revert) each fail. So that every colour sits
   where this reads it, no element carries a style or an opacity attribute and no stylesheet is
   linked or imported. Every colour the workbook sets on a cell, a border, a tab or a chart is a
   token's RGB: a theme, indexed, preset or automatic colour fails, but for exactly one, Excel's
   default text colour (theme 1), which the font of every cell the workbook leaves unstyled carries.
   Every hex literal in the four export modules is a token value.
3. Grounds and contrast: the rules that paint behind text or change how its colours composite
   (background, box-shadow, border-image, opacity, filter, backdrop-filter, mix-blend-mode, mask) are
   exactly `.masthead`'s background and the marked block's (`blockquote.located`, the same tint) in each
   PDF, so a new ground fails until it is named here. The gate cannot know where a selector matches, so
   every text colour the stylesheet sets clears 4.5:1 on every ground the stylesheet paints and on the
   white page: the masthead's document kind, the table header and the tertiary ink among them. Every
   workbook cell's font clears 4.5:1 on its fill, or on white without one.
4. Sage (D-041, applying D-030 to the exports): the wordmark's "Nerd" is the only text that takes a
   brand colour. A PDF link rests in its line's ink with an underline, since print has no pointer to
   turn it sage and a link's text can carry an index ("Source filing 1"): each PDF's stylesheet
   carries that rule, which also keeps WeasyPrint's default link blue off the page, and no rule that
   styles a link sets a colour other than inherit, currentColor or an ink, or drops the underline (a
   text decoration without it, a zero thickness, `all`). A rule styles a link when its subject, read
   with cssselect2 (WeasyPrint's own selector parser), can select an `<a href>`: an `a`, a link
   pseudo-class (`:link`, `:any-link`, `:visited`, `:local-link`), an `[href]` attribute, `:is()` or
   `:where()` holding one, or a subject with no type, class, id or other attribute (`*`,
   `:first-child`, `:not(p)`), with a pseudo-element or without; and when it selects an `<a>` in
   either rendered PDF, whatever its form. In the workbook no text takes sage (the sheet tabs keep it).
5. The wordmark (the PDF rendition of BG-27): "Earnings" in the ink, "Nerd" italic in brand-strong,
   from a real italic. WeasyPrint never slants a face it lacks, and with only Inter's roman declared
   "Nerd" and every quote printed upright. So the wordmark's family declares an italic face, every
   @font-face names a vendored file (one that exists in backend/app/assets/fonts, so a system font
   declared under a brand family's name fails), no two faces share a family and a style (WeasyPrint keys
   an embedded font by family, style and weight, never by file, so a second face merges with the first and
   one file's glyph ids print from the other's program), and whose italic bit (OS/2 fsSelection bit 0, read with
   fontTools) matches the style it is declared as, and in each rendered PDF every font WeasyPrint
   embeds for text the stylesheet sets italic (its descriptor's Italic flag) is an italic program.
6. Every glyph prints in a vendored face (WP3-A): the export inspection found 52 daggers, 2 arrows and 614
   list bullets set in DejaVu, since the latin subsets of Inter and Geist Mono lack †, → and ◦. Inter's
   symbols faces carry them, roman and italic, and "≈" (the 20-F and 6-K prompts ask for "≈US$X" beside a
   convenience translation) and the won and rupee signs; the data face's stack falls back to Inter before
   the system's monospace, and a nested list takes the disc (•, in Inter's latin subset) instead of
   WeasyPrint's circle. Over both fixtures, an analysis whose narrative (roman and an em), grid and count
   line print the marks (its nested operand list prints the disc), and a 20-F summary that prints "≈", "₩"
   and "₹" in prose, a figure column and a quote, each with a mark, a space and a mark, the two fixtures
   printing every mark in MARKS between them: WeasyPrint logs no .notdef;
   every font each render embeds, whatever kind of program, is of a vendored face's family (Inter Var, Inter
   Var Symbols or Geist Mono) and encoded by Identity-H; every glyph a composite font's text map pairs with
   text is, in a file declared for the font's family and style, that text's glyph or one its GSUB reaches
   from it (WeasyPrint keeps the files' glyph ids), so a font that merges two files matches neither; and
   every marker of a nested list is the disc. The glyph check reads the text as given and as NFC (HarfBuzz
   composes a decomposed cluster where the face encodes the composed letter, "E" and U+0308 to Geist Mono's
   Edieresis, and where it does not, GSUB may ligate the parts, "C" and U+0301 to Inter's Cacute), accepts a
   space the face does not encode drawn as the face's own space (HarfBuzz's fallback), and follows ligatures in
   chains ("<->" in Inter: "->" to arrowright, then less and arrowright to arrowboth); a reduction that runs
   past _REDUCE_STATES states fails the test rather than stall it. What it cannot see: a substitute GSUB reaches that a different feature setting
   would not have chosen, a shaping HarfBuzz does outside GSUB other than those two, which it would flag
   though the render is right, and text a font draws outside its text map (WeasyPrint keeps the first text
   each glyph draws).

The planted forms in test_the_checks_read_what_they_claim score each reader of the source (stylesheet,
attributes, faces, workbook, module literals) on the forms its rule can take and on look-alikes that
must pass; the render check scores its reader on a planted PDF of its own. Out of reach (review): a colour written in an export module
other than as a hex, where the fixtures do not render it; a rule that reaches a link only through a
class or id the fixtures' links do not carry; the colours Excel and WeasyPrint draw where nothing
sets one (a chart's text); and the embedded brand PNG's pixels.
"""
from __future__ import annotations

import ast
import functools
import io
import logging
import re
import unicodedata
import zipfile
from pathlib import Path
from urllib.parse import unquote, urlparse

import cssselect2
import pytest
import tinycss2
import tinycss2.color4
import tinyhtml5
import yaml
from cssselect2 import parser as selectors
from fontTools.ttLib import TTFont
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Color, Font, PatternFill, Side

import tests.unit.test_excel_export as workbook_fixtures
import tests.unit.test_export_service as export_fixtures
from app.services.excel_export_service import CHART_SERIES, build_analysis_workbook
from app.services.export_service import ExportService
from app.services.pdf_branding import PALETTE, render_branded_pdf

ROOT = Path(__file__).resolve().parents[3]
EXPORT_MODULES = ("export_service.py", "excel_export_service.py", "analysis_export_metadata.py", "pdf_branding.py")
_FRONTMATTER = re.match(r"^---\n(.*?)\n---\n", (ROOT / "DESIGN.md").read_text(encoding="utf-8"), re.DOTALL)
TOKENS = {name: str(value).upper() for name, value in yaml.safe_load(_FRONTMATTER.group(1))["colors"].items()
          if str(value).startswith("#")}
TOKEN_VALUES = set(TOKENS.values())
SAGE = {value for name, value in TOKENS.items() if name.startswith("brand")}
WHITE = TOKENS["white"]
INKS = {TOKENS[name] for name in ("text-primary-light", "text-secondary-light", "text-tertiary-light")}

# Each PALETTE key's token, by role.
ROLES = {
    "brand": "brand",
    "brand_strong": "brand-strong",
    "brand_weak": "brand-weak",
    "brand_border": "brand-border",
    "background": "background-light",
    "panel": "panel-light",
    "ink": "text-primary-light",
    "ink_secondary": "text-secondary-light",
    "ink_tertiary": "text-tertiary-light",
    "border": "border-light",
}
# A hex in a module literal: 3, 4, 6 or 8 digits, the alpha forms read as written (never a token).
HEX = re.compile(r"#(?:[0-9A-Fa-f]{8}|[0-9A-Fa-f]{6}|[0-9A-Fa-f]{4}|[0-9A-Fa-f]{3})(?![0-9A-Za-z])")
BARE_HEX = re.compile(r"^(?:[0-9A-Fa-f]{2})?[0-9A-Fa-f]{6}$")  # openpyxl's RGB or aRGB: no "#"
HEX_COLOUR = re.compile(r"^#(?:[0-9A-Fa-f]{3}|[0-9A-Fa-f]{6})$")  # the only hex a token can be

# What can name a colour in a CSS value besides a hex, and what reaches one the gate cannot read.
COLOUR_FUNCTIONS = {
    "rgb", "rgba", "hsl", "hsla", "hwb", "lab", "lch", "oklab", "oklch", "color", "color-mix",
    "device-cmyk", "light-dark", "contrast-color",
}
INDIRECTIONS = {"var", "env", "attr"}
SYSTEM_COLOURS = {
    "accentcolor", "accentcolortext", "activetext", "buttonborder", "buttonface", "buttontext", "canvas",
    "canvastext", "field", "fieldtext", "graytext", "highlight", "highlighttext", "linktext", "mark",
    "marktext", "selecteditem", "selecteditemtext", "visitedtext", "activeborder", "activecaption",
    "appworkspace", "background", "buttonhighlight", "buttonshadow", "captiontext", "inactiveborder",
    "inactivecaption", "inactivecaptiontext", "infobackground", "infotext", "menu", "menutext", "scrollbar",
    "threeddarkshadow", "threedface", "threedhighlight", "threedlightshadow", "threedshadow", "window",
    "windowframe", "windowtext",
}
CSS_WIDE = {"initial", "unset", "revert", "revert-layer"}  # inherit is the one that keeps a read colour
CARRIES_COLOUR = re.compile(
    r"colou?r|^(?:-[a-z]+-)?(?:all|background|border|outline|text-decoration|text-emphasis|column-rule|fill"
    r"|stroke|box-shadow|text-shadow)"
)
# A rule that paints behind text or changes how its colours composite.
GROUND = re.compile(
    r"^(?:-[a-z]+-)?(?:background|box-shadow$|border-image|opacity$|filter$|backdrop-filter$|mix-blend-mode$|mask)"
)
SVG_PAINT = {"fill", "stroke", "color", "stop-color", "flood-color", "lighting-color"}
OPACITY_ATTRIBUTES = {"opacity", "fill-opacity", "stroke-opacity", "stop-opacity", "flood-opacity"}
LINK_PSEUDO_CLASSES = {"link", "any-link", "visited", "local-link"}
WORKBOOK_DEFAULT_TEXT = "theme 1"  # Excel's default text colour, on every cell the workbook leaves unstyled


def normal(hex_value: str) -> str:
    digits = hex_value.lstrip("#").upper()
    return "#" + ("".join(c * 2 for c in digits) if len(digits) == 3 else digits[-6:])


def luminance(hex_value: str) -> float:
    def channel(c: int) -> float:
        s = c / 255
        return s / 12.92 if s <= 0.03928 else ((s + 0.055) / 1.055) ** 2.4
    r, g, b = (int(normal(hex_value)[i:i + 2], 16) for i in (1, 3, 5))
    return 0.2126 * channel(r) + 0.7152 * channel(g) + 0.0722 * channel(b)


def contrast(ink: str, ground: str) -> float:
    lighter, darker = sorted((luminance(ink), luminance(ground)), reverse=True)
    return (lighter + 0.05) / (darker + 0.05)


def selector_parts(selector_list: str) -> list[str]:
    """A selector list's selectors, split at its top-level commas (a comma inside `:is()` stays)."""
    parts: list[list] = [[]]
    for token in tinycss2.parse_component_value_list(selector_list):
        if token.type == "literal" and token.value == ",":
            parts.append([])
        else:
            parts[-1].append(token)
    return [text for text in (tinycss2.serialize(part).strip() for part in parts) if text]


def css_rules(css: str) -> list[tuple[str, list[tuple[str, str]]]]:
    """(name, declarations) for every block of a stylesheet, as WeasyPrint reads it. A style rule is named
    by its selectors, a rule nested in another under its full selector (`&` replaced, else a descendant of
    its parent), a rule inside @media or @supports by its own; an at-rule's own block by the at-rule, so a
    margin box nested in @page reads "@page @top-left". Every declaration, in order, none collapsed."""
    rules: list[tuple[str, list[tuple[str, str]]]] = []

    def contents(node) -> list:
        return tinycss2.parse_blocks_contents(node.content, skip_comments=True, skip_whitespace=True)

    def block(name: str, items: list, parents: list[str] | None) -> None:
        declarations = [(item.lower_name, tinycss2.serialize(item.value).strip())
                        for item in items if item.type == "declaration"]
        if declarations:
            rules.append((name, declarations))
        for item in items:
            if item.type == "qualified-rule":
                own = selector_parts(tinycss2.serialize(item.prelude))
                full = own if parents is None else [
                    part.replace("&", parent) if "&" in part else f"{parent} {part}" for parent in parents for part in own
                ]
                block(", ".join(full), contents(item), full)
            elif item.type == "at-rule" and item.content is not None:
                at = f"@{item.lower_at_keyword} {tinycss2.serialize(item.prelude).strip()}".strip()
                # In a style rule, a conditional group's declarations are that rule's (CSS nesting).
                inner = ", ".join(parents) if parents is not None else f"{name} {at}" if name.startswith("@") else at
                block(inner, contents(item), parents)

    block("", tinycss2.parse_stylesheet(css, skip_comments=True, skip_whitespace=True), None)
    return rules


def last(rules: list[tuple[str, list[tuple[str, str]]]], name: str) -> dict[str, str]:
    """The declarations in force in the last block of that name (the last of each property wins)."""
    return dict([block for rule, block in rules if rule == name][-1])


def stylesheet(html: str) -> str:
    return "".join(re.findall(r"<style\b[^>]*>(.*?)</style>", html, re.DOTALL))


def elements(html: str) -> list:
    """Every element of the document as WeasyPrint parses it (tinyhtml5), the inline SVG's included."""
    return list(tinyhtml5.parse(html).iter())


def colours_in(name: str, value: str) -> list[str]:
    """Every colour a declaration names, as written: a hex, a colour function, a colour keyword (named,
    system, `transparent`, `currentColor`), an indirection (var(), env(), attr()), and on a property that
    carries a colour, a CSS-wide keyword other than inherit, whose colour is not one the gate can read."""
    found: list[str] = []
    nodes = tinycss2.parse_component_value_list(value, skip_comments=True)

    def walk(items) -> None:
        for node in items:
            if node.type == "hash":
                found.append("#" + node.value)
            elif node.type == "function":
                if node.lower_name in COLOUR_FUNCTIONS | INDIRECTIONS:
                    found.append(tinycss2.serialize([node]))
                else:
                    walk(node.arguments)
            elif node.type in ("() block", "[] block", "{} block"):
                walk(node.content)
            elif node.type == "ident" and (
                node.lower_value in SYSTEM_COLOURS or tinycss2.color4.parse_color(node) is not None
            ):
                found.append(node.value)

    walk(nodes)
    words = [node for node in nodes if node.type not in ("whitespace", "comment")]
    if len(words) == 1 and words[0].type == "ident" and words[0].lower_value in CSS_WIDE and CARRIES_COLOUR.search(name):
        found.append(words[0].value)
    return found


def is_token(colour_value: str) -> bool:
    return HEX_COLOUR.match(colour_value) is not None and normal(colour_value) in TOKEN_VALUES


def pdf_colour_problems(html: str) -> list[str]:
    """Each colour a PDF carries that is not a token's hex or currentColor, and each place a colour could
    sit unread: a style or opacity attribute, a linked or imported stylesheet."""
    found = []
    for rule, block in css_rules(stylesheet(html)):
        for name, value in block:
            found += [f"{rule} {{ {name}: {value} }} names {c}" for c in colours_in(name, value)
                      if not is_token(c) and c.lower() != "currentcolor"]
    found += [f"@import {tinycss2.serialize(node.prelude).strip()}"
              for node in tinycss2.parse_stylesheet(stylesheet(html), skip_comments=True, skip_whitespace=True)
              if node.type == "at-rule" and node.lower_at_keyword == "import"]
    for element in elements(html):
        tag = element.tag.rsplit("}", 1)[-1]
        for attribute, value in element.attrib.items():
            if attribute == "style":
                found.append(f"<{tag} style={value!r}>: a style attribute")
            elif attribute in OPACITY_ATTRIBUTES:
                found.append(f"<{tag} {attribute}={value!r}>: an opacity attribute")
            elif attribute in SVG_PAINT:
                found += [f"<{tag} {attribute}={value!r}> names {c}" for c in colours_in(attribute, value)
                          if not is_token(c) and c.lower() != "currentcolor"]
        if tag == "link" and "stylesheet" in element.get("rel", "").lower().split():
            found.append(f"<link rel=stylesheet href={element.get('href')!r}>: a linked stylesheet")
    return found


def text_colours(html: str) -> list[tuple[str, str]]:
    """(rule, text colour) for every hex a `color` declaration sets (any other colour fails above)."""
    return [(rule, normal(value)) for rule, block in css_rules(stylesheet(html))
            for name, value in block if name == "color" and HEX_COLOUR.match(value)]


def ground_rules(html: str) -> dict[str, list[str]]:
    """Each rule that paints behind text or changes how its colours composite, with the properties that do."""
    found: dict[str, set[str]] = {}
    for rule, block in css_rules(stylesheet(html)):
        for name, _ in block:
            if GROUND.match(name):
                found.setdefault(rule, set()).add(name)
    return {rule: sorted(names) for rule, names in found.items()}


def grounds(html: str) -> set[str]:
    """The white page and every hex a rule paints behind text."""
    return {WHITE} | {normal(c) for _, block in css_rules(stylesheet(html)) for name, value in block
                      if GROUND.match(name) for c in colours_in(name, value) if HEX_COLOUR.match(c)}


def low_contrast(html: str) -> list[str]:
    """Each text colour that falls below 4.5:1 on a ground the stylesheet paints, or on the white page."""
    return [f"{rule} {ink} on {ground} is {contrast(ink, ground):.2f}:1"
            for rule, ink in text_colours(html) for ground in sorted(grounds(html)) if contrast(ink, ground) < 4.5]


def _subject(selector) -> list:
    """The simple selectors of a parsed selector's subject: its last compound."""
    tree = selector.parsed_tree
    while isinstance(tree, selectors.CombinedSelector):
        tree = tree.right
    return tree.simple_selectors


def _holds_for_every_link(simple: list) -> bool:
    """Whether a compound selects every <a href> (`a`, `:link`, `:any-link`, `[href]`), so `:not()` of it
    selects none."""
    return bool(simple) and all(
        (isinstance(s, selectors.LocalNameSelector) and s.lower_local_name == "a")
        or (isinstance(s, selectors.PseudoClassSelector) and s.name in ("link", "any-link"))
        or (isinstance(s, selectors.AttributeSelector) and s.lower_name == "href" and s.operator is None)
        for s in simple
    )


def _styles_a_link(simple: list) -> bool:
    """Whether a compound can select an <a href>. It cannot with a type other than `a`, a `:not()` of
    every link, or an `:is()` or `:where()` none of whose members can. It can when an `a`, a link
    pseudo-class, an [href] or such an `:is()` marks it, whatever class narrows it, and when it is bare
    (no type, class, id or other attribute: `*`, `:first-child`, `:not(p)`), since that selects a link too."""
    logical = (selectors.MatchesAnySelector, selectors.SpecificityAdjustmentSelector)
    if any(isinstance(s, selectors.LocalNameSelector) and s.lower_local_name != "a" for s in simple):
        return False
    if any(isinstance(s, selectors.NegationSelector) and any(_holds_for_every_link(_subject(m)) for m in s.selector_list)
           for s in simple):
        return False
    if any(isinstance(s, logical) and not any(_styles_a_link(_subject(m)) for m in s.selector_list) for s in simple):
        return False
    marked = any(
        isinstance(s, (selectors.LocalNameSelector, *logical))
        or (isinstance(s, selectors.PseudoClassSelector) and s.name in LINK_PSEUDO_CLASSES)
        or (isinstance(s, selectors.AttributeSelector) and s.lower_name == "href")
        for s in simple
    )
    bare = not any(
        isinstance(s, (selectors.ClassSelector, selectors.IDSelector))
        or (isinstance(s, selectors.AttributeSelector) and s.lower_name != "href")
        for s in simple
    )
    return marked or bare


def link_rules(html: str) -> list[tuple[str, list[tuple[str, str]]]]:
    """(selector, declarations) for each selector that styles a link: by its form (its subject can select
    an <a href>, a link pseudo-class, an [href] or a bare subject among them) or by what it selects in this
    rendered PDF (an <a>). A selector WeasyPrint cannot compile styles nothing, as WeasyPrint drops it."""
    root = cssselect2.ElementWrapper.from_html_root(tinyhtml5.parse(html))
    anchors = [element for element in root.iter_subtree() if element.local_name == "a"]
    found = []
    for rule, block in css_rules(stylesheet(html)):
        if rule.startswith("@"):
            continue
        for part in selector_parts(rule):
            try:
                parsed, compiled = list(selectors.parse(part)), cssselect2.compile_selector_list(part)
            except cssselect2.SelectorError:
                continue
            if any(_styles_a_link(_subject(selector)) for selector in parsed) or any(
                selector.test(anchor) for selector in compiled for anchor in anchors
            ):
                found.append((part, block))
    return found


def link_problems(html: str) -> list[str]:
    """Each declaration a rule styling a link makes against D-041: a colour other than inherit,
    currentColor or an ink, or a text decoration that drops the underline."""
    found = []
    for part, block in link_rules(html):
        for name, value in block:
            colours = colours_in(name, value)
            if name == "color" and not (value.lower() in ("inherit", "currentcolor")
                                        or (HEX_COLOUR.match(value) and normal(value) in INKS)):
                found.append(f"{part} sets {name}: {value}")
            elif name != "color" and any(
                c.lower() != "currentcolor" and not (HEX_COLOUR.match(c) and normal(c) in INKS) for c in colours
            ):
                found.append(f"{part} sets {name}: {value}")
            elif re.fullmatch(r"(?:-[a-z]+-)?text-decoration(?:-line)?", name) and "underline" not in value.lower():
                found.append(f"{part} sets {name}: {value}")
            elif re.fullmatch(r"(?:-[a-z]+-)?text-decoration-thickness", name) and re.fullmatch(
                r"[+-]?0*\.?0*[a-z%]*", value.lower()
            ):
                found.append(f"{part} sets {name}: {value}")
            elif name == "all":
                found.append(f"{part} sets {name}: {value}")
    return found


def font_faces(html: str) -> list[dict]:
    """Each @font-face of a PDF's stylesheet: its family, its declared style and the file its src names."""
    faces = []
    for rule, block in css_rules(stylesheet(html)):
        if rule == "@font-face":
            declared = dict(block)
            source = re.search(r"url\(\s*[\"']?([^\"')]+)", declared.get("src", ""))
            faces.append({
                "family": declared.get("font-family", "").strip("\"' "),
                "style": declared.get("font-style", "normal").lower(),
                "file": Path(unquote(urlparse(source.group(1)).path)) if source else None,
            })
    return faces


def is_italic_program(font: str | Path | bytes) -> bool:
    """A font program's own italic bit: OS/2 fsSelection bit 0."""
    return bool(TTFont(io.BytesIO(font) if isinstance(font, bytes) else font)["OS/2"].fsSelection & 1)


# Where every face a PDF declares lives: the vendored files, never a system font under a brand family's name.
VENDORED_FONT_DIR = ROOT / "backend" / "app" / "assets" / "fonts"


def font_face_problems(html: str, italic_family: str) -> list[str]:
    """Each way a PDF's faces fail: the italic family declares no italic face, two faces share a family and a
    style, a face's file is not a vendored one (it resolves outside backend/app/assets/fonts, through a
    symlink or a ".." too), a face's file is missing, or a file's italic bit disagrees with the style it is
    declared as."""
    faces = font_faces(html)
    found = [] if any(f["family"] == italic_family and f["style"] == "italic" for f in faces) else [
        f"'{italic_family}' declares no italic face"]
    declared: dict[tuple[str, str], list[str]] = {}
    for face in faces:
        declared.setdefault((face["family"], face["style"]), []).append(face["file"].name if face["file"] else "no src")
    found += [f"'{family}' {style} is declared by {len(files)} files ({', '.join(files)}): WeasyPrint merges them"
              for (family, style), files in declared.items() if len(files) > 1]
    for face in faces:
        name = f"'{face['family']}' {face['style']} ({face['file'].name if face['file'] else 'no src'})"
        if face["file"] is None:
            found.append(f"{name} names no file")
        elif not face["file"].resolve().is_relative_to(VENDORED_FONT_DIR.resolve()):
            found.append(f"{name} is not a vendored file in backend/app/assets/fonts")
        elif not face["file"].is_file():
            found.append(f"{name} names no file")
        elif is_italic_program(face["file"]) != (face["style"] in ("italic", "oblique")):
            found.append(f"{name} is declared {face['style']} but its file's italic bit says otherwise")
    return found


def embedded_fonts(pdf: bytes) -> list[tuple[str, bool, bool]]:
    """(name, asked for italic, italic program) for each font an uncompressed WeasyPrint PDF embeds: its
    descriptor's Italic flag (bit 7, set from the style the text asked for) and the embedded program's
    own italic bit."""
    objects = dict(re.findall(rb"\n(\d+) 0 obj\n(.*?)\nendobj", pdf, re.DOTALL))
    found = []
    for body in objects.values():
        descriptor = re.search(rb"/Type /FontDescriptor/FontName /([^\s/]+).*?/Flags (\d+).*?/FontFile2 (\d+) 0 R",
                               body, re.DOTALL)
        if descriptor:
            name, flags, program = descriptor.groups()
            stream = objects[program]
            data = stream[stream.index(b"stream\n") + len(b"stream\n"):stream.rindex(b"\nendstream")]
            found.append((name.decode(), bool(int(flags) & 64), is_italic_program(data)))
    return found


def _pdfs() -> dict[str, str]:
    analysis = export_fixtures.TestAnalysisPdfHtml()
    summary, filing = export_fixtures._make_summary_and_filing(export_fixtures._full_sections())
    return {
        "filing summary": ExportService().generate_pdf_html(summary, filing),
        "analysis": ExportService().generate_analysis_pdf_html(analysis._analysis(), analysis._company()),
    }


def _workbook_bytes() -> bytes:
    return build_analysis_workbook(workbook_fixtures._dataset(), exported_at=workbook_fixtures.EXPORTED_AT)


def _xl(color) -> str:
    """An openpyxl colour as the gate reads it: "#RRGGBB" for an RGB, its kind and index otherwise."""
    if color is None:
        return "automatic"
    if color.type == "rgb":
        return normal(color.rgb)
    if color.type == "theme":
        return f"theme {color.theme}" + (f" tint {color.tint}" if color.tint else "")
    if color.type == "indexed":
        return f"indexed {color.indexed}"
    return str(color.type)


def workbook_colours(data: bytes) -> list[tuple[str, str, str | None]]:
    """(where, colour, ground) for every colour the workbook sets, whatever its kind; the ground is the
    cell's fill (or white) for a font, None for anything else."""
    found: list[tuple[str, str, str | None]] = []
    book = load_workbook(io.BytesIO(data))
    for sheet in book.worksheets:
        tab = sheet.sheet_properties.tabColor
        if tab is not None:
            found.append((f"{sheet.title} tab", _xl(tab), None))
        for row in sheet.iter_rows():
            for cell in row:
                where = f"{sheet.title}!{cell.coordinate}"
                fill = _xl(cell.fill.start_color) if cell.fill.patternType == "solid" else None
                if fill:
                    found.append((where + " fill", fill, None))
                for side in ("left", "right", "top", "bottom"):
                    edge = getattr(cell.border, side)
                    if edge is not None and edge.style is not None:
                        found.append((f"{where} {side} border", _xl(edge.color), None))
                if cell.value not in (None, ""):
                    found.append((where + " text", _xl(cell.font.color), fill if fill and fill.startswith("#") else WHITE))
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        for name in archive.namelist():
            if name.startswith("xl/charts/"):
                xml = archive.read(name).decode("utf-8")
                found += [(name, normal(value), None) for value in re.findall(r'srgbClr val="([0-9A-Fa-f]{6})"', xml)]
                found += [(name, f"{kind}{attributes}", None)
                          for kind, attributes in re.findall(r"<a:(schemeClr|sysClr|prstClr|scrgbClr|hslClr)\b([^>]*)>", xml)]
    return found


def workbook_off_token(found: list[tuple[str, str, str | None]]) -> list[tuple[str, str]]:
    return [(where, value) for where, value, ground in found
            if value not in TOKEN_VALUES and not (where.endswith(" text") and value == WORKBOOK_DEFAULT_TEXT)]


def hex_literals(source: str) -> list[tuple[int, str]]:
    """(line, colour) for each hex in a string literal of Python source: "#RRGGBB" for a 3- or 6-digit hex
    and for openpyxl's RGB or aRGB ("FF1A1A17"), a CSS hex with alpha as written (never a token)."""
    found = []
    for node in ast.walk(ast.parse(source)):
        if isinstance(node, ast.Constant) and isinstance(node.value, str):
            values = [normal(value) if HEX_COLOUR.match(value) else value for value in HEX.findall(node.value)]
            if BARE_HEX.match(node.value):
                values.append(normal(node.value))  # an aRGB keeps its colour in the last six digits
            found += [(node.lineno, value) for value in values]
    return sorted(found)


def module_hex_literals() -> list[tuple[str, int, str]]:
    return [(module, line, value) for module in EXPORT_MODULES
            for line, value in hex_literals((ROOT / "backend" / "app" / "services" / module).read_text(encoding="utf-8"))]


def test_each_palette_key_holds_its_roles_token():
    assert set(PALETTE) == set(ROLES), "name the new key's token in ROLES (or drop the key)"
    drift = {key: (PALETTE[key], ROLES[key], TOKENS[ROLES[key]]) for key in PALETTE
             if normal(PALETTE[key]) != TOKENS[ROLES[key]]}
    assert drift == {}, "PALETTE key: (value, its role's token, that token's value)"
    assert [normal(value) for value in CHART_SERIES] == [TOKENS[f"chart-{n}"] for n in range(1, 7)]


def test_every_rendered_and_written_colour_is_a_token():
    off_token = [(name, problem) for name, html in _pdfs().items() for problem in pdf_colour_problems(html)]
    off_token += workbook_off_token(workbook_colours(_workbook_bytes()))
    off_token += [(f"{module}:{line}", value) for module, line, value in module_hex_literals() if value not in TOKEN_VALUES]
    assert off_token == []
    # The readers found what they read: the stylesheets' hexes, the mark's fill and the workbook's colours.
    assert all(sum(len(colours_in(n, v)) for _, b in css_rules(stylesheet(html)) for n, v in b) > 20 for html in _pdfs().values())
    assert len(workbook_colours(_workbook_bytes())) > 100


def test_the_grounds_are_the_page_the_masthead_and_the_marked_block_alone():
    """A rule that paints behind text is a ground every text colour must clear; the masthead's tint and the
    marked block's (a located excerpt, D-044) are the only ones, one token, so a new ground (a tinted section,
    an inset shadow, a page colour) fails here first."""
    for name, html in _pdfs().items():
        assert ground_rules(html) == {".masthead": ["background"], "blockquote.located": ["background"]}, name
        assert grounds(html) == {WHITE, TOKENS["brand-weak"]}, name


def test_every_text_colour_clears_its_ground():
    checked, low = set(), []
    for name, html in _pdfs().items():
        checked |= {rule for rule, _ in text_colours(html)}
        low += [f"{name}: {problem}" for problem in low_contrast(html)]
    for where, ink, ground in workbook_colours(_workbook_bytes()):
        if ground is not None and ink.startswith("#") and contrast(ink, ground) < 4.5:
            low.append(f"workbook: {where} {ink} on {ground} is {contrast(ink, ground):.2f}:1")
    assert low == []
    # The roles the rule names are among the rules read, so a pass covers them.
    assert {".masthead .doc-kind", "th", ".meta", ".footnote", "@page @top-left"} <= checked


def test_sage_colours_only_the_wordmarks_nerd():
    """D-041: a header, a label, a figure and a link rest in the ink; the wordmark's "Nerd" is the only
    sage text in the exports."""
    sage_text = []
    for name, html in _pdfs().items():
        for rule, ink in text_colours(html):
            if ink in SAGE:
                sage_text += [(name, part) for part in selector_parts(rule) if part != ".masthead .wordmark em"]
    sage_text += [(where, ink) for where, ink, ground in workbook_colours(_workbook_bytes())
                  if ground is not None and ink in SAGE]
    assert sage_text == []


def test_every_pdf_link_rests_in_the_ink_with_an_underline():
    """Print has no pointer, so a link never turns sage: it keeps its line's ink and an underline in
    the tertiary ink. Without the rule, WeasyPrint's default stylesheet would paint links blue."""
    for name, html in _pdfs().items():
        rules = css_rules(stylesheet(html))
        assert any(rule == "a" for rule, _ in rules), f"the {name} PDF's stylesheet carries no rule for its links"
        links = last(rules, "a")
        assert links.get("color") == "inherit" and "underline" in links.get("text-decoration", ""), name
        assert normal(links.get("text-decoration-color", "")) == TOKENS["text-tertiary-light"], name
        assert link_problems(html) == [], name  # no scoped rule turns a link sage or drops its underline
        assert "a" in [part for part, _ in link_rules(html)], name  # the reader found the link rule
    assert '<a href="' in _pdfs()["analysis"]  # the fixture renders links (its data-note letters)


def _wordmark_family(html: str) -> str:
    return last(css_rules(stylesheet(html)), ".masthead .wordmark")["font-family"].split(",")[0].strip("\"' ")


def test_the_wordmark_reads_earnings_in_the_ink_and_nerd_italic_in_brand_strong():
    for name, html in _pdfs().items():
        assert '<span class="wordmark">Earnings<em>Nerd</em></span>' in html, name
        rules = css_rules(stylesheet(html))
        assert normal(last(rules, ".masthead .wordmark")["color"]) == TOKENS["text-primary-light"], name
        nerd = last(rules, ".masthead .wordmark em")
        assert normal(nerd["color"]) == TOKENS["brand-strong"] and nerd["font-style"] == "italic", name
        # The italic is a real one: the family declares an italic face, and each face's file is what it says.
        assert _wordmark_family(html) == "Inter Var", name
        assert font_face_problems(html, _wordmark_family(html)) == [], name
        assert len(font_faces(html)) >= 3, name  # the reader found the faces (Inter's two, Geist Mono)


def test_every_font_asked_for_italic_embeds_an_italic_program():
    """The PDF a reader gets: every font WeasyPrint embeds for text the stylesheet sets italic (the wordmark's
    "Nerd", a quote, an em, and a mark the symbols faces carry in either) is an italic program, never the roman
    under an italic name."""
    weasyprint = pytest.importorskip("weasyprint")
    htmls = {**_pdfs(), "analysis with every mark": _every_mark_analysis(), "20-F summary with every mark": _every_mark_summary()}
    try:
        pdfs = {name: weasyprint.HTML(string=html).write_pdf(uncompressed_pdf=True) for name, html in htmls.items()}
    except Exception as exc:  # missing system libs in odd envs: not what this test reads
        pytest.skip(f"WeasyPrint runtime unavailable: {exc}")
    for name, pdf in pdfs.items():
        fonts = embedded_fonts(pdf)
        assert [font for font, asked, italic in fonts if asked and not italic] == [], name
        assert any(asked for _, asked, _ in fonts), name  # the fixture prints italic text: the wordmark's "Nerd"
    # Planted: the same PDF with Inter's roman alone, which printed "Nerd" upright under an italic name.
    summary = _pdfs()["filing summary"]
    roman_only = re.sub(r"@font-face \{[^}]*font-style: italic;\s*\}", "", summary, count=1)
    assert roman_only != summary, "the planted PDF still declares the italic face"
    upright = [font for font, asked, italic in embedded_fonts(weasyprint.HTML(string=roman_only).write_pdf(uncompressed_pdf=True))
               if asked and not italic]
    assert upright and all(font.endswith("Inter-Var-Oblique") for font in upright), upright


# WeasyPrint writes the family a run resolved to as its font descriptor's /FontFamily: the vendored faces'.
VENDORED_FAMILIES = frozenset({"Inter Var", "Inter Var Symbols", "Geist Mono"})
# The marks the latin subsets lack, which Inter's symbols faces carry: the every-mark fixtures must print each
# (the glyph gate reads them), and test_vendored_font_sources.py holds pdf_branding.SYMBOLS_RANGE equal to it.
MARKS = frozenset("†‡←→◦≈₩₹")
_SYMBOLS_FACE = re.compile(r"@font-face \{[^}]*Inter-Variable-(?:Italic-)?symbols\.woff2[^}]*\}")  # roman and italic


def embedded_font_families(pdf: bytes) -> list[str]:
    """The family of every font an uncompressed WeasyPrint PDF embeds, whatever kind of program it is
    (embedded_fonts reads TrueType programs alone): each font descriptor's /FontFamily, a literal or a
    UTF-16 hex string, and "Type 3" for each Type 3 font, which carries none."""
    families = []
    for literal, hexed in re.findall(rb"/Type /FontDescriptor/FontName /[^/\s]+/FontFamily (?:\(([^)]*)\)|<([0-9A-Fa-f]*)>)", pdf):
        families.append(bytes.fromhex(hexed.decode()).decode("utf-16") if hexed else literal.decode("latin-1"))
    assert len(families) == pdf.count(b"/Type /FontDescriptor"), "a font descriptor this reader did not read"
    return families + ["Type 3"] * len(re.findall(rb"/Subtype /Type3\b", pdf))


@functools.lru_cache(maxsize=None)
def vendored_glyphs(path: Path) -> tuple[list[str], dict[int, str], dict[str, frozenset[str]],
                                         dict[tuple[str, ...], frozenset[str]]]:
    """A vendored face's glyph order, its cmap, each glyph's closure under its GSUB substitutions (single,
    multiple, alternate and reverse chaining, through extension lookups; contextual lookups call these, so a
    glyph any feature can turn it into, itself included) and each ligature's glyphs by its components."""
    font = TTFont(path)
    order, cmap = font.getGlyphOrder(), font.getBestCmap()
    edges: dict[str, set[str]] = {}
    ligatures: dict[tuple[str, ...], set[str]] = {}
    for lookup in font["GSUB"].table.LookupList.Lookup if "GSUB" in font else ():
        for table in lookup.SubTable:
            kind = lookup.LookupType
            if kind == 7:
                kind, table = table.ExtensionLookupType, table.ExtSubTable
            if kind == 1:
                pairs = [(source, [target]) for source, target in table.mapping.items()]
            elif kind in (2, 3):
                pairs = list((table.mapping if kind == 2 else table.alternates).items())
            elif kind == 8:
                pairs = [(source, [target]) for source, target in zip(table.Coverage.glyphs, table.Substitute)]
            elif kind == 4:
                for first, entries in table.ligatures.items():
                    for entry in entries:
                        ligatures.setdefault((first, *entry.Component), set()).add(entry.LigGlyph)
                continue
            else:
                continue
            for source, targets in pairs:
                edges.setdefault(source, set()).update(targets)
    closure = {}
    for glyph in order:
        seen, todo = {glyph}, [glyph]
        while todo:
            for target in edges.get(todo.pop(), ()):
                if target not in seen:
                    seen.add(target)
                    todo.append(target)
        closure[glyph] = frozenset(seen)
    return order, cmap, closure, {key: frozenset(value) for key, value in ligatures.items()}


# The search visits every mix of each position's substitutes before it rejects a cluster, so a long wrong one
# grows as their product: Geist Mono's ">>>>>" ran past 20 s, and now fails at the cap in about 0.3 s. The
# fixtures' renders take at most 18 states.
_REDUCE_STATES = 50_000


def _reduces_to(sequence: tuple[str, ...], glyph: str, closure: dict[str, frozenset[str]],
                ligatures: dict[tuple[str, ...], frozenset[str]]) -> bool:
    """Whether GSUB can turn the glyph ``sequence`` into ``glyph``: by substituting any glyph with one its
    lookups reach and ligating any run of adjacent glyphs, in any order and as often as it takes ("<->": calt
    ligates "->" to arrowright, then less and arrowright to arrowboth). A search past _REDUCE_STATES states
    fails, naming the cluster, rather than stall the test."""
    seen, todo = {sequence}, [sequence]
    while todo:
        current = todo.pop()
        if current == (glyph,):
            return True
        assert len(seen) <= _REDUCE_STATES, f"reducing {sequence} to {glyph} passed {_REDUCE_STATES} states"

        steps = [current[:index] + (form,) + current[index + 1:]
                 for index, member in enumerate(current) for form in closure.get(member, ())]
        steps += [current[:start] + (ligature,) + current[end:]
                  for start in range(len(current)) for end in range(start + 2, len(current) + 1)
                  for ligature in ligatures.get(current[start:end], ())]
        for step in steps:
            if step not in seen:
                seen.add(step)
                todo.append(step)
    return False


def _entry_problem(path: Path, gid: int, text: str) -> str | None:
    """Why ``gid`` is not a glyph the vendored file at ``path`` draws for ``text``, or None when it is: what
    GSUB makes of the text's glyphs (a substitute, a ligature, a chain of them), the text read as given or as
    NFC. HarfBuzz's normaliser composes a decomposed cluster where the face encodes the composed letter ("E"
    and U+0308 to Geist Mono's Edieresis); where it does not, the parts stay and GSUB may ligate them ("C" and
    U+0301 to Inter's Cacute, by its ccmp). A space the face does not encode (a thin, narrow or ideographic
    space) HarfBuzz draws with the face's own space, or a substitute of it (Inter's tabular space)."""
    order, cmap, closure, ligatures = vendored_glyphs(path)
    if gid >= len(order):
        return f"glyph {gid} is outside its {len(order)} glyphs"
    glyph, composed = order[gid], unicodedata.normalize("NFC", text)
    if len(composed) == 1 and unicodedata.category(composed) == "Zs" and ord(composed) not in cmap and 0x20 in cmap:
        return None if glyph in closure[cmap[0x20]] else f"it draws the unencoded {composed!r} as {glyph}, not its space"
    readings = [tuple(cmap.get(ord(char)) for char in reading) for reading in dict.fromkeys((composed, text))]
    if any(None not in bases and _reduces_to(bases, glyph, closure, ligatures) for bases in readings):
        return None
    if all(None in bases for bases in readings):
        return f"it encodes no {composed!r} (drawn as {glyph})"
    return f"it draws {text!r} as {next(bases for bases in readings if None not in bases)[0]}, not {glyph}"


def glyph_problems(html: str, pdf: bytes) -> list[str]:
    """Each way the text of an uncompressed WeasyPrint render of ``html`` is not drawn with its own glyphs.

    WeasyPrint subsets with the original glyph ids (retain_gids), so a composite font's ToUnicode map pairs a
    glyph id of the file it came from with the text drawn by it. Each font must match a file ``html`` declares
    for its family and its style (the descriptor's italic flag; a family with no face in that style falls back
    to its normal one): every pair is that file's glyph for the text, or a substitute or ligature its GSUB
    reaches. Two files declared as one family and style become one embedded font (WeasyPrint keys fonts by
    family, style and weight), and then neither file matches: one file's glyph ids index the other's program,
    whichever registered first. A font encoded otherwise than by Identity-H is one WeasyPrint drew a .notdef
    for, a character no face in its stack carries. A family no vendored face declares is the family gate's."""
    faces = font_faces(html)
    objects = dict(re.findall(rb"\n(\d+) 0 obj\n(.*?)\nendobj", pdf, re.DOTALL))

    def stream(number: bytes) -> bytes:
        body = objects[number]
        return body[body.index(b"stream\n") + len(b"stream\n"):body.rindex(b"\nendstream")]

    found = []
    composites = [body for body in objects.values() if re.search(rb"/Subtype /Type0\b", body)]
    assert composites, "the render embeds no composite font this reader found"
    for body in composites:
        name = re.search(rb"/BaseFont /([^/\s]+)", body)[1].decode().split("+")[-1]
        encoding = re.search(rb"/Encoding (/[^\s/>\]]+|\d+ 0 R)", body)[1].decode()
        if encoding != "/Identity-H":
            found.append(f"{name} is encoded by {encoding}, not /Identity-H: WeasyPrint drew a .notdef glyph in it")
            continue
        descendant = objects[re.search(rb"/DescendantFonts \[(\d+) 0 R\]", body)[1]]
        descriptor = objects[re.search(rb"/FontDescriptor (\d+) 0 R", descendant)[1]]
        literal, hexed = re.search(rb"/FontFamily (?:\(([^)]*)\)|<([0-9A-Fa-f]*)>)", descriptor).groups()
        family = bytes.fromhex(hexed.decode()).decode("utf-16") if hexed else literal.decode("latin-1")
        style = "italic" if int(re.search(rb"/Flags (\d+)", descriptor)[1]) & 64 else "normal"
        files = ([face["file"] for face in faces if face["family"] == family and face["style"] == style]
                 or [face["file"] for face in faces if face["family"] == family and face["style"] == "normal"])
        if not files:
            continue
        pairs = []
        for block in re.findall(r"beginbfchar(.*?)endbfchar", stream(re.search(rb"/ToUnicode (\d+) 0 R", body)[1]).decode("latin-1"), re.DOTALL):
            pairs += [(int(gid, 16), bytes.fromhex(code).decode("utf-16-be")) for gid, code in re.findall(r"<([0-9A-Fa-f]+)>\s*<([0-9A-Fa-f]+)>", block)]
        assert pairs, f"{name}'s text map holds no glyph this reader found"
        misses = {path.name: [miss for gid, text in pairs if text and (miss := _entry_problem(path, gid, text))] for path in files}
        if all(misses.values()):
            found.append(f"{name} ({family} {style}) matches no file declared for it: "
                         + "; ".join(f"{file}: {', '.join(miss[:3])}" for file, miss in misses.items()))
    return found


def rendered_text(pdf: bytes) -> set[str]:
    """Every character an uncompressed WeasyPrint PDF's composite fonts pair with a glyph in their text maps."""
    found: set[str] = set()
    for text_map in re.findall(rb"beginbfchar(.*?)endbfchar", pdf, re.DOTALL):
        found |= {char for _, code in re.findall(rb"<([0-9A-Fa-f]+)>\s*<([0-9A-Fa-f]+)>", text_map)
                  for char in bytes.fromhex(code.decode()).decode("utf-16-be")}
    return found


def rendered(html: str) -> tuple[bytes, list[str]]:
    """An uncompressed render of ``html`` and the .notdef warnings WeasyPrint logged while drawing it (a glyph
    no font in the run's stack has, drawn as an empty box)."""
    import weasyprint

    messages: list[str] = []
    handler = logging.Handler()
    handler.emit = lambda record: messages.append(record.getMessage())
    logger = logging.getLogger("weasyprint")
    logger.addHandler(handler)
    try:
        pdf = weasyprint.HTML(string=html).write_pdf(uncompressed_pdf=True)
    finally:
        logger.removeHandler(handler)
    return pdf, [message for message in messages if ".notdef" in message]


def foreign_fonts(html: str) -> list[str]:
    """The families an uncompressed render of ``html`` embeds that no vendored face declares."""
    import weasyprint

    families = embedded_font_families(weasyprint.HTML(string=html).write_pdf(uncompressed_pdf=True))
    assert families, "the render embeds no font this reader found"
    return sorted(set(families) - VENDORED_FAMILIES)


def list_markers(html: str) -> dict[int, set[str]]:
    """Each list marker's text in a render of ``html``, by how many lists hold its item (1: a top-level list)."""
    import weasyprint
    from weasyprint.formatting_structure import boxes

    found: dict[int, set[str]] = {}

    def walk(box, lists: tuple[int, ...]) -> None:
        if getattr(box, "element_tag", None) in ("ul", "ol") and id(box.element) not in lists:
            lists = (*lists, id(box.element))
        if isinstance(box, boxes.TextBox) and box.element_tag == "li::marker":
            found.setdefault(len(lists), set()).add(box.text)
        for child in getattr(box, "children", ()):  # an absolute placeholder hands on its box's children
            walk(child, lists)

    for page in weasyprint.HTML(string=html).render().pages:
        walk(page._page_box, ())
    return found


def _every_mark_analysis() -> str:
    """The analysis PDF with each mark the latin subsets lack in reach: an arrow in the narrative (the trend
    detectors join a move with " → "), a calculated figure, whose dagger follows it in the grid's data face
    and in the data notes' count line, and an operand with an operand of its own, a list in the methods list."""
    analysis = export_fixtures.TestAnalysisPdfHtml()._analysis()
    # A mark, a space and a mark keep the space in the marks' run (Pango), in roman and in italic (an em). The
    # thin space comes first: WeasyPrint's text map keeps the first text a glyph draws, so the reader reads the
    # face's space as drawn for it (HarfBuzz's fallback for a space the face does not encode).
    analysis.narrative_md = ("Net margin moved 20.0% → 24.3% [2]. Marked †\u2009‡ and † ‡. Back ← and ◦ open."
                             " *Moved 20.0% → 24.3%, † ‡.*")
    analysis.dataset_json["series"][0]["points"][2]["provenance"] = {
        "method": "calculated", "validation": "passed", "reasons": [], "formula": "annual_minus_ytd",
        "inputs": [{"concept": "revenue", "value": 2500, "unit": "USD", "provenance": {
            "method": "calculated", "validation": "passed", "formula": "sum",
            "inputs": [{"concept": "product_revenue", "value": 2000, "unit": "USD"}]}}],
    }
    html = ExportService().generate_analysis_pdf_html(analysis, export_fixtures.TestAnalysisPdfHtml()._company())
    assert "→" in html and '$1,500 †</td>' in html and "calculated †" in html and '<ul class="operands">' in html
    assert "<em>Moved 20.0% → 24.3%, † ‡.</em>" in html and "Marked †\u2009‡ and † ‡. Back ← and ◦ open." in html
    return html


def _every_mark_summary() -> str:
    """A 20-F summary PDF with each mark a foreign filer's summary can print beyond the latin subsets: "≈US$X"
    beside a convenience translation, which the 20-F and 6-K prompts ask for, in prose and in the metrics
    table's figure column (the data face), and the won and rupee signs before a figure; a mark, a space and a
    mark (Pango keeps the space in the marks' run) in prose, in a figure cell and in a quote, set italic."""
    sections = export_fixtures._full_sections()
    sections["executive_snapshot"]["key_points"] += [
        "Revenue of RMB 941.2B (≈US$129B, convenience translation)",
        "Korea ₩1.2T → ₩1.4T; India ₹4,100 → ₹4,500 crore; RMB 941.2B → ≈US$129B; ≈ ₩1.2T; † ‡; → ₩",
    ]
    # The first space the data face draws in a figure cell is a narrow no-break space, which Geist Mono does not
    # encode: HarfBuzz draws its own space, and the reader reads that pairing first.
    sections["financial_highlights"]["table"].insert(0, {
        "metric": "Shares outstanding", "current_period": "1\u202f234M", "prior_period": "1\u202f200M",
        "change": "+2.8%", "commentary": "A narrow no-break space in the data face."})
    sections["financial_highlights"]["table"] += [
        {"metric": "Revenue in US dollars", "current_period": "≈US$129B", "prior_period": "≈US$120B",
         "change": "+7.5%", "commentary": "Convenience translation at the stated rate."},
        {"metric": "Korea revenue", "current_period": "₩1.2T → ₩1.4T", "prior_period": "† ‡",
         "change": "+16.7%", "commentary": "A mark, a space and a mark in the data face."},
    ]
    sections["segment_performance"].append(
        {"segment": "Korea", "revenue": "₩1.2T", "change": "-3%", "commentary": "India revenue of ₹4,100 crore"})
    sections["management_discussion_insights"]["quotes"].append(
        {"speaker": "CFO", "quote": "Korea moved ₩1.2T → ₩1.4T, ≈US$1B; † ‡."})
    summary, filing = export_fixtures._make_summary_and_filing(sections)
    filing.filing_type = "20-F"
    html = ExportService().generate_pdf_html(summary, filing)
    assert "(≈US$129B," in html and '<td class="num">≈US$129B</td>' in html and "₩1.2T" in html and "₹4,100" in html
    assert '<td class="num">₩1.2T → ₩1.4T</td><td class="num">† ‡</td>' in html
    assert '<td class="num">1\u202f234M</td>' in html
    assert "<blockquote>“Korea moved ₩1.2T → ₩1.4T, ≈US$1B; † ‡.”" in html
    return html


def test_every_glyph_prints_in_a_vendored_face(tmp_path):
    pytest.importorskip("weasyprint")
    pdfs = {**_pdfs(), "analysis with every mark": _every_mark_analysis(), "20-F summary with every mark": _every_mark_summary()}
    try:
        renders = {name: rendered(html) for name, html in pdfs.items()}
    except Exception as exc:  # missing system libs in odd envs: not what this test reads
        pytest.skip(f"WeasyPrint runtime unavailable: {exc}")
    assert {name: notdef for name, (_, notdef) in renders.items()} == {name: [] for name in pdfs}, "a .notdef glyph"
    families = {name: embedded_font_families(pdf) for name, (pdf, _) in renders.items()}
    assert all(families.values()), families  # the reader found the fonts
    assert {name: sorted(set(found) - VENDORED_FAMILIES) for name, found in families.items()} == {name: [] for name in pdfs}
    assert {name: glyph_problems(pdfs[name], pdf) for name, (pdf, _) in renders.items()} == {name: [] for name in pdfs}
    drawn = set().union(*(rendered_text(renders[name][0]) for name in ("analysis with every mark", "20-F summary with every mark")))
    assert MARKS <= drawn, f"the every-mark fixtures print none of {sorted(MARKS - drawn)}"
    for name in ("analysis with every mark", "20-F summary with every mark"):  # the reader saw the marks' faces
        assert {"Inter Var Symbols"} <= set(families[name]), name
        assert any(font.endswith("Inter-Var-Symbols-Italic") and asked and italic
                   for font, asked, italic in embedded_fonts(renders[name][0])), name
    markers = list_markers(pdfs["analysis with every mark"])
    assert markers[2] == {"• "} and "• " in markers[1], markers  # the operands list sits in the methods list

    # What HarfBuzz makes of a cluster, each a right render the reader must pass: a ligature chain ("<->" in
    # Inter), a decomposed letter Geist Mono encodes composed ("E" and U+0308 to Edieresis, read as NFC) and one
    # Inter's latin cut encodes only as parts ("C" and U+0301, which its ccmp ligates to Cacute, read as given).
    for body in ("<p>A &lt;-&gt; B</p>", '<table><tr><td class="num">E\u0308 1</td></tr></table>', "<p>C\u0301racovia</p>"):
        shaped = render_branded_pdf(title="Shaped", doc_kind="Shaped", meta_html="", body_html=body)
        pdf, notdef = rendered(shaped)
        assert (notdef, glyph_problems(shaped, pdf)) == ([], []), body

    # Planted: the same PDF without the symbols face sets the dagger and the arrow in DejaVu; with the data
    # face's stack as it was, the grid's dagger takes the system's monospace; without the disc rule a nested
    # list prints WeasyPrint's circle.
    every_mark = pdfs["analysis with every mark"]
    no_symbols = _SYMBOLS_FACE.sub("", every_mark)
    old_stack = every_mark.replace("'Geist Mono', 'Inter Var', 'Inter Var Symbols', monospace", "'Geist Mono', monospace")
    no_disc = every_mark.replace("ul ul { list-style-type: disc; }", "")
    assert no_symbols != every_mark and old_stack != every_mark and no_disc != every_mark
    assert foreign_fonts(no_symbols) and foreign_fonts(old_stack), "a fallback face must read as foreign"
    summary = pdfs["20-F summary with every mark"]
    assert foreign_fonts(_SYMBOLS_FACE.sub("", summary)), "≈, ₩ and ₹ need the symbols faces"
    assert foreign_fonts(no_disc) == [] and list_markers(no_disc)[2] == {"◦ "}  # the face covers ◦: only the marker reads it

    # Planted: the symbols face declared as a second 'Inter Var' normal face, its first form, merges with the
    # latin one, whichever registers first. Latin first: the analysis grid drew "†" with a latin glyph, the
    # 20-F summary "₩" as "A". Symbols first (a mark is the first text at its size): "ALIBABA" from the
    # symbols program.
    merge = ("font-family: 'Inter Var Symbols';", "font-family: 'Inter Var';")
    for name in ("analysis with every mark", "20-F summary with every mark"):
        merged = pdfs[name].replace(*merge, 1)
        assert merged != pdfs[name] and glyph_problems(merged, rendered(merged)[0]), name
    symbols_first = render_branded_pdf(title="Planted", doc_kind="Planted", meta_html="",
                                       body_html="<p>→ ₩</p><p>ALIBABA GROUP</p>").replace(*merge, 1)
    assert glyph_problems(symbols_first, rendered(symbols_first)[0])
    # Planted: the symbols cut without its two spaces, the first cuts' form. A space between two marks stays in
    # their run, and WeasyPrint draws it as .notdef and encodes that font by a CMap of its own.
    from fontTools import subset

    roman = VENDORED_FONT_DIR / "Inter-Variable-symbols.woff2"
    spaceless = TTFont(roman)
    subsetter = subset.Subsetter(subset.Options(layout_features=["*"], notdef_outline=True))
    subsetter.populate(unicodes=[code for code in spaceless.getBestCmap() if code not in (0x20, 0xA0)])
    subsetter.subset(spaceless)
    spaceless.flavor = None
    spaceless.save(tmp_path / "symbols-without-spaces.ttf")
    planted = summary.replace(f"url('{roman.as_uri()}') format('woff2')",
                              f"url('{(tmp_path / 'symbols-without-spaces.ttf').as_uri()}') format('truetype')")
    assert planted != summary
    pdf, notdef = rendered(planted)
    assert notdef and any("not /Identity-H" in problem for problem in glyph_problems(planted, pdf)), notdef


PLANTED_STYLES = """
.masthead { background: #ECF2EE; } .masthead .doc-kind { color: #6B7280; } th { color: #3C6650; }
@page { size: A4; @top-left { color: #636A77; } }
"""


def test_a_runaway_reduction_fails_rather_than_stalls():
    """The reducer tries every mix of each position's substitutes before it rejects a cluster, so a long wrong one
    grows as their product: Geist Mono's ">>>>>" (23 substitutes a place) ran past 20 s. It fails at
    _REDUCE_STATES instead, naming the cluster, in well under a second."""
    _, cmap, closure, ligatures = vendored_glyphs(VENDORED_FONT_DIR / "GeistMono-latin.woff2")
    runaway = tuple(cmap[ord(char)] for char in ">>>>>")
    with pytest.raises(AssertionError, match=rf"reducing \('greater', .*\) to \.notdef passed {_REDUCE_STATES} states"):
        _reduces_to(runaway, ".notdef", closure, ligatures)


def test_the_checks_read_what_they_claim():
    """Planted forms: each reader names the forms its rule can take and passes the look-alikes."""
    assert round(contrast("#6B7280", TOKENS["brand-weak"]), 2) == 4.26  # the old tertiary on the tint
    assert round(contrast(TOKENS["text-tertiary-light"], TOKENS["brand-weak"]), 2) == 4.80
    page = f"<style>{PLANTED_STYLES}</style>"
    assert text_colours(page) == [(".masthead .doc-kind", "#6B7280"), ("th", "#3C6650"), ("@page @top-left", "#636A77")]
    assert low_contrast(page) == [".masthead .doc-kind #6B7280 on #ECF2EE is 4.26:1"]

    # 2. Colour values: every form but a token's hex, inherit and currentColor fails, wherever it sits.
    flagged = {
        ".ref { color: rgb(60, 102, 80); }": "rgb(60, 102, 80)",
        ".sources strong { color: green; }": "green",
        ".x { color: hsl(150 26% 32%); }": "hsl(150 26% 32%)",
        ".x { color: #3C6650FF; }": "#3C6650FF",
        ".x { color: #3C68; }": "#3C68",
        ".x { color: transparent; }": "transparent",
        ".x { color: var(--sage); }": "var(--sage)",
        ".x { color: color-mix(in srgb, #3C6650, #FFFFFF); }": "color-mix(in srgb, #3C6650, #FFFFFF)",
        ".x { color: CanvasText; }": "CanvasText",
        ".x { color: initial; }": "initial",
        ".x { color: revert; }": "revert",
        ".x { border-bottom: 1px solid rebeccapurple; }": "rebeccapurple",
        ".x { text-decoration: underline green; }": "green",
        ".x { box-shadow: 0 0 0 1px rgb(0 0 0 / 50%); }": "rgb(0 0 0 / 50%)",
        ".sources { .x { color: green; } }": "green",  # a nested rule
        "@media print { .ref { color: green; } }": "green",  # a rule in a conditional group
        ".x { color: #636A77; color: green; }": "green",  # a declaration a later one overrides
    }
    for css, form in flagged.items():
        problems = pdf_colour_problems(f"<style>{css}</style>")
        assert len(problems) == 1 and problems[0].endswith(f"names {form}"), (css, problems)
    passing = (
        "a { color: inherit; } .x { color: currentColor; border: 0; } .y { color: #636A77; }"
        " th { text-transform: uppercase; font-family: 'Geist Mono', monospace; } ul { list-style: none; }"
        ' @page { @bottom-right { content: "Page " counter(page); color: #636A77; } } a { text-decoration: underline; }'
    )
    assert pdf_colour_problems(f"<style>{passing}</style>") == []
    assert pdf_colour_problems('<svg><path fill="#3C6650" d="M0"/><path fill="none" stroke="currentColor" d="M0"/></svg>') == []
    assert pdf_colour_problems(
        '<div style="color: #636A77">x</div><svg><path fill="green" d="M0"/><path fill-opacity="0.5" d="M0"/></svg>'
        '<link rel="stylesheet" href="x.css"><style>@import url(x.css);</style><style media="print">.r { color: red; }</style>'
    ) == [
        ".r { color: red } names red",  # a <style> with attributes is read too
        "@import url(x.css)",
        "<div style='color: #636A77'>: a style attribute",
        "<path fill='green'> names green",
        "<path fill-opacity='0.5'>: an opacity attribute",
        "<link rel=stylesheet href='x.css'>: a linked stylesheet",
    ]

    # A module literal's hex: the alpha forms fail, an aRGB reads as its colour.
    assert hex_literals('A = "#3C6650FF"\nB = "FF1A1A17"\nC = f"#3c6650 {x}"\nD = "#3C68"') == [
        (1, "#3C6650FF"), (2, "#1A1A17"), (3, "#3C6650"), (4, "#3C68"),
    ]
    # The workbook: a theme, an indexed and an automatic colour fail; an unstyled cell's default text passes.
    book = Workbook()
    sheet = book.active
    sheet.title = "Planted"
    sheet["A1"] = "unstyled"
    sheet["A2"], sheet["A2"].font = "theme", Font(color=Color(theme=4))
    sheet["A3"], sheet["A3"].fill = "indexed", PatternFill("solid", start_color=Color(indexed=10))
    sheet["A4"].border = Border(bottom=Side(style="thin"))
    buffer = io.BytesIO()
    book.save(buffer)
    assert workbook_off_token(workbook_colours(buffer.getvalue())) == [
        ("Planted!A2 text", "theme 4"), ("Planted!A3 fill", "indexed 10"), ("Planted!A4 bottom border", "automatic"),
    ]

    # 5. Faces: a family with no italic face, a face whose file is the roman and a missing file each fail.
    fonts = ROOT / "backend" / "app" / "assets" / "fonts"
    roman, italic = (fonts / "Inter-Variable-latin.woff2").as_uri(), (fonts / "Inter-Variable-Italic-latin.woff2").as_uri()
    symbols = (fonts / "Inter-Variable-symbols.woff2").as_uri()
    face = "@font-face {{ font-family: 'Inter Var'; src: url('{}') format('woff2'); font-style: {}; }}"
    assert font_face_problems(f"<style>{face.format(roman, 'normal')}</style>", "Inter Var") == [
        "'Inter Var' declares no italic face"]
    assert font_face_problems(f"<style>{face.format(roman, 'normal')} {face.format(roman, 'italic')}</style>", "Inter Var") == [
        "'Inter Var' italic (Inter-Variable-latin.woff2) is declared italic but its file's italic bit says otherwise"]
    assert font_face_problems(f"<style>{face.format(italic, 'normal')} {face.format(italic, 'italic')}</style>", "Inter Var") == [
        "'Inter Var' normal (Inter-Variable-Italic-latin.woff2) is declared normal but its file's italic bit says otherwise"]
    assert font_face_problems(f"<style>{face.format(fonts.as_uri() + '/missing.woff2', 'italic')}</style>", "Inter Var") == [
        "'Inter Var' italic (missing.woff2) names no file"]
    assert font_face_problems(f"<style>{face.format(roman, 'normal')} {face.format(italic, 'italic')}</style>", "Inter Var") == []
    assert font_face_problems(f"<style>{face.format(roman, 'normal')} {face.format(symbols, 'normal')} {face.format(italic, 'italic')}</style>",
                              "Inter Var") == [
        "'Inter Var' normal is declared by 2 files (Inter-Variable-latin.woff2, Inter-Variable-symbols.woff2): WeasyPrint merges them"]
    # A system font declared under a brand family's name fails whether or not this machine has it, and so does
    # a path that leaves the vendored directory through "..".
    system = "file:///usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    escaped = fonts.as_uri() + "/../brand/earningsnerd-mark.png"
    assert font_face_problems(
        f"<style>{face.format(system, 'normal')} {face.format(italic, 'italic')} "
        f"{face.format(escaped, 'normal').replace('Inter Var', 'Geist Mono')}</style>", "Inter Var") == [
        "'Inter Var' normal (DejaVuSans.ttf) is not a vendored file in backend/app/assets/fonts",
        "'Geist Mono' normal (earningsnerd-mark.png) is not a vendored file in backend/app/assets/fonts",
    ]

    # 3. Grounds: a rule painting behind text joins the grounds, and every text colour is held to each.
    tinted = f"<style>{PLANTED_STYLES} .methods {{ background: #E5E7EB; }} .footnote {{ color: #636A77; }}</style>"
    assert ground_rules(tinted) == {".masthead": ["background"], ".methods": ["background"]}
    assert ".footnote #636A77 on #E5E7EB is 4.40:1" in low_contrast(tinted)
    for css in (".x { box-shadow: inset 0 0 0 99px #E5E7EB; }", "@page { background: #ECF2EE; }",
                ".x { opacity: 0.8; }", ".x { background-image: linear-gradient(#ECF2EE, #ECF2EE); }",
                ".x { .y { background-color: #ECF2EE; } }", ".x { filter: brightness(0.9); }"):
        assert len(ground_rules(f"<style>{css}</style>")) == 1, css

    # 4. Links: each form that styles a link is read, and each look-alike is not.
    links = (
        "a { color: inherit; text-decoration: underline; } .sources a { text-decoration: none; }"
        " .methods > a:link { color: #3C6650; } .x a.y { text-decoration-color: #3C6650; }"
        " .data-note a { color: blue; } .ok a { color: #374151; } abbr { color: #3C6650; }"
        " .sources :any-link { text-decoration: none; } .methods [href] { color: green; }"
        " .x :link { color: #3C6650; } .x * { text-decoration-line: none; } .x :is(a, b) { color: #4F7A63; }"
        " .x :not(p) { text-decoration: none; } .x a::after { color: #3C6650; } .x :visited { all: unset; }"
        " .x { a { text-decoration: none; } } .x a { text-decoration-thickness: 0; } .src { text-decoration: none; }"
        " .footnote { text-decoration: none; } .x :not(a) { color: #3C6650; } .x :not(:link) { color: #3C6650; }"
        " td.num { color: #374151; }"
    )
    assert link_problems(f'<style>{links}</style><p><a href="#n" class="src">1</a></p>') == [
        ".sources a sets text-decoration: none", ".methods > a:link sets color: #3C6650",
        ".x a.y sets text-decoration-color: #3C6650", ".data-note a sets color: blue",
        ".sources :any-link sets text-decoration: none", ".methods [href] sets color: green",
        ".x :link sets color: #3C6650", ".x * sets text-decoration-line: none", ".x :is(a, b) sets color: #4F7A63",
        ".x :not(p) sets text-decoration: none", ".x a::after sets color: #3C6650", ".x :visited sets all: unset",
        ".x a sets text-decoration: none", ".x a sets text-decoration-thickness: 0",
        ".src sets text-decoration: none",  # a class the rendered link carries: read by what it selects
    ]
